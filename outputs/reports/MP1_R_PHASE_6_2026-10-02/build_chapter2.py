"""Reproducible Phase 6 draft build; no retrieval, adjudication, or Phase 7 work."""
from pathlib import Path
import json, re, hashlib, zipfile
from docx import Document
from docx.shared import Cm, Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[3]
RUN = Path(__file__).parent
DATE = '2026-10-02'
DOCX = f'docs/project/MP1_R_CHAPTER_2_EVIDENCE_CONTROLLED_DRAFT_{DATE}.docx'
TRACE = f'outputs/plans/MP1_R_PHASE_6_CHAPTER_2_CLAIM_TRACEABILITY_{DATE}.json'
REOPEN = f'outputs/plans/MP1_R_PHASE_6_SOURCE_REOPEN_LOG_{DATE}.json'
P5PATH = f'outputs/plans/MP1_R_PHASE_5_PHASE_6_HANDOFF_{DATE}.json'
p5 = json.loads((ROOT / P5PATH).read_text())
assert p5['phase5_status'] == 'READY_FOR_PHASE_6_CHAPTER_2_DRAFTING'
manifest = json.loads((RUN / 'input_manifest.json').read_text())
assert all(hashlib.sha256((ROOT / x['path']).read_bytes()).hexdigest() == x['sha256'] for x in manifest)

# Resume the preserved prose inventory, not a historical Chapter 2 narrative.
ns = {'__name__': 'preserved_phase6_inventory'}
script = (RUN / 'checkpoint_backup/generate_mp1_ch2_pre_resume.py').read_text()
exec(script.split('# Build DOCX.')[0], ns)
items = [(k, v) for k, v in ns['items'] if k != 'ref']
legacy_order = ns['ref_order']
legacy_inverse = {v: k for k, v in legacy_order.items()}

def replace_paragraph(prefix, text):
    hits = [i for i, (k, v) in enumerate(items) if k == 'p' and v[0].startswith(prefix)]
    assert len(hits) == 1, prefix
    items[hits[0]] = ('p', (text, []))

replace_paragraph('Chương này xây dựng',
    'Chương này tổng hợp cơ sở khoa học cho việc đặc trưng hóa ứng xử uốn phụ thuộc áp suất của một bó dây TiNi có cấu hình xác định. Hai thành phần cần được phân biệt là kết cấu chịu uốn và nguồn áp suất dẫn động bằng hợp kim nhớ hình dạng (shape memory alloy, SMA). Phạm vi tổng quan gồm cơ chế giam giữ, tiếp xúc và ma sát, đặc tính nhiệt–cơ của TiNi, cùng các mô hình có thể dùng làm đối chứng. Nguồn lò xo SMA–piston được xem xét ở mức nhiệm vụ hỗ trợ và tích hợp; đáp ứng kết cấu vẫn có thể được khảo sát bằng nguồn áp suất bên ngoài.')
replace_paragraph('Mạch lập luận được tổ chức',
    'Tổng quan được tổ chức theo quan hệ giữa kích thích, đường truyền lực và đại lượng quan sát. Các kết quả thực nghiệm được diễn giải trong phạm vi vật liệu, hình học, tải và điều kiện biên của từng nghiên cứu. Mô hình hoặc phép thử trên hệ tương tự cung cấp cơ sở tham chiếu, nhưng không thay thế cho kiểm chứng trên bó TiNi. Cách tiếp cận này giúp phân biệt điều đã được thiết lập trong các kiến trúc liên quan với vấn đề cần nhận dạng hoặc đặc trưng hóa riêng cho MP1-R.')
replace_paragraph('Phần cuối chương tổng hợp',
    'Phần cuối chương kết nối chuỗi từ áp suất đến đáp ứng uốn và xác định giới hạn suy luận tại mỗi mắt xích. Trọng tâm là ánh xạ từ phản lực hướng tâm đến lực pháp tuyến tiếp xúc (contact normal force) cục bộ tại tiếp diện dây–dây (wire–wire) và dây–ống bao (wire–sleeve). Từ đó, chương dẫn đến một câu hỏi nghiên cứu chính và hai câu hỏi phụ, đồng thời giữ các cách giải thích vật liệu, tiếp xúc thông thường và tương tác bổ sung như những khả năng cần phân biệt bằng quan sát.')
replace_paragraph('Trong các kết cấu biến đổi',
    'Trong các kết cấu biến đổi độ cứng, giam giữ có thể hạn chế chuyển động tương đối của nhiều phần tử thông qua bao mềm, môi trường hạt hoặc lực ép cơ học. Ứng xử dính–trượt (stick–slip) thay đổi đáp ứng chịu tải của toàn cụm, nhưng đường truyền tải pháp tuyến phụ thuộc cách bố trí từng kiến trúc [1], [3], [4]. Vì vậy, sự thay đổi lực–độ võng hoặc mô men–độ cong không tự nó xác định lực tiếp xúc cục bộ hay tiếp diện chi phối.')
replace_paragraph('Tổng quan của Caro',
    'Caro và Carmichael phân loại các cơ chế thay đổi độ cứng của kết cấu giam giữ nhiều lớp, gồm ma sát, cản trở chuyển động bằng hình học và những cơ chế khác [2]. Phạm vi bài tổng quan tập trung vào cấu trúc lớp; nó không bao quát các hệ hạt hoặc sợi. Nguồn này hữu ích để định vị cơ chế và thuật ngữ, còn các mệnh đề về thí nghiệm hay mô hình cụ thể cần dựa trên nghiên cứu gốc. Sự phân biệt này tránh dùng một tổng quan lớp làm bằng chứng lặp lại độc lập cho bó dây.')
replace_paragraph('Liu và cộng sự trình bày',
    'Các thí nghiệm trên hệ hạt của Liu và cộng sự bổ sung một đường giam giữ khác với bao hút chân không của Bai và cộng sự: áp suất dương được dùng để điều chỉnh đáp ứng của cụm hạt–túi [3]. Đối chiếu hai hệ cho thấy áp suất dương và hút chân không đều có thể đi kèm thay đổi ứng xử uốn, nhưng vật liệu phần tử, hướng phản lực và độ mềm bao khác nhau [1], [3]. Bằng chứng vì vậy xác lập hiệu ứng trong từng kiến trúc, chứ không xác lập một quan hệ áp suất–độ cứng phổ quát cho bó TiNi.')
replace_paragraph('Zhang và Yao xây dựng',
    'So với môi trường hạt, chuỗi sợi nylon chịu áp suất dương của Zhang và Yao gần hơn với MP1-R về câu hỏi cơ học nhiều phần tử. Công trình kết hợp mô hình các nhánh dính, trượt một phần và trượt hoàn toàn với thí nghiệm uốn; độ nhạy theo áp suất khác nhau giữa các nhánh [4]. Tuy nhiên, chuỗi dùng sợi nylon, các phần liên kết cứng và túi bơm phồng, trong khi MP1-R dùng bó TiNi với đường phản lực và tiếp diện chưa được đồng nhất. Một phần tham số được nhận dạng từ đáp ứng tổng thể. Do đó, cách phân nhánh là cơ sở tham chiếu phù hợp, còn hệ số truyền áp suất và lực tiếp xúc không được chuyển trực tiếp sang cấu hình mới.')
replace_paragraph('Zhang, Yao, Zhao',
    'Mô hình liên tục của Zhang, Yao, Zhao và Wei xét dầm nhiều lớp polyvinyl chloride (PVC) dưới hút chân không, thay vì sợi chịu áp suất dương [5]. Mô hình dùng ngưỡng ma sát để mô tả vùng dính–trượt và được đối chiếu với phần tử hữu hạn cùng đáp ứng toàn cục của mẫu. Hạn chế liên quan đến ứng suất pháp tuyến ngang, hiệu ứng đầu mút và vùng trượt ở biến dạng lớn cho thấy mức đơn giản hóa đường tiếp xúc ảnh hưởng khả năng dự đoán. Công trình này cung cấp tương tự lý thuyết–cơ học cho MP1-R; nó không xác nhận bó TiNi, giam giữ sợi áp suất dương hoặc lực pháp tuyến cục bộ của hệ đề xuất.')
replace_paragraph('Tại một tiếp diện',
    'Trong mô hình ma sát kiểu Coulomb, độ lớn lực tiếp tuyến bị giới hạn bởi tích của lực pháp tuyến với hệ số ma sát. Cách biểu diễn này được dùng trong mô hình sợi và dầm nhiều lớp đã khảo sát [4], [5]. Đối với một tiếp diện quy ước i, giới hạn tối thiểu được viết là:')
for i, (k, v) in enumerate(items):
    if k == 'eq': items[i] = ('eq', '|Fₜ,ᵢ| ≤ μᵢ Nᵢ')
replace_paragraph('Trong đó F_t',
    'Trong đó Fₜ,ᵢ là lực tiếp tuyến tại tiếp diện i, Nᵢ là lực pháp tuyến và μᵢ là hệ số ma sát hiệu dụng của tiếp diện đó. Đẳng thức ở biên giới hạn chỉ mô tả điều kiện trượt trong giả thiết mô hình, không có nghĩa lực ma sát luôn bằng μᵢNᵢ. Áp suất chất lưu p, có đơn vị lực trên diện tích, cũng không đồng nhất với Nᵢ, có đơn vị lực. Việc chuyển từ p sang Nᵢ đòi hỏi đường phản lực, diện tích truyền tải và phân bố tiếp xúc phù hợp; đây là phần chưa được thiết lập cho bó TiNi [3]–[5].')
replace_paragraph('Barsi, Carboni',
    'Ở dây cáp thép ngắn, Barsi, Carboni và Lacarbonara xây dựng mô hình cơ học và đối chiếu các giới hạn độ cứng bằng thí nghiệm uốn [8]. Khác với các mô hình áp suất của sợi và lớp, đường tải ở đây chịu ảnh hưởng của hình học cáp và điều kiện đầu mút, còn trạng thái trượt được dùng để xác định các giới hạn đáp ứng [4], [5], [8]. Kết quả tạo cơ sở cho một mô hình đối chứng về độ cứng, nhưng không cung cấp luật đầy đủ cho vòng trễ hoặc phân bố áp suất–tiếp xúc của MP1-R.')
replace_paragraph('Vahidi và cộng sự trình bày',
    'Vahidi và cộng sự kết hợp mô hình vật liệu SMA với tiếp xúc ma sát trong mô phỏng cáp xoắn đơn và xoắn kép [9]. Công trình chứng minh rằng chuyển pha và tiếp xúc có thể cùng được đưa vào một mô hình số thông thường, thay vì mặc nhiên cần một luật tương tác mới. Tuy nhiên, đáp ứng mô phỏng thuộc miền kéo của cáp xoắn, không phải uốn ngang của bó chịu áp suất gần mẫu; nó cũng không phải một thí nghiệm riêng xác nhận bó TiNi của MP1-R. Khả năng dự đoán trong cấu hình đề xuất vẫn cần được kiểm tra.')
replace_paragraph('Các đại lượng đo cũng',
    'Độ cứng uốn (bending stiffness) cần được khai báo bằng một đại lượng vận hành cụ thể. Độ dốc lực–độ võng kF = dF/dδ có đơn vị N/m, trong đó F là lực ngoài và δ là độ võng; tỷ số F/δ là độ cứng cát tuyến, không luôn bằng độ dốc. Với mô men M và độ cong κ, dM/dκ là độ cứng tiếp tuyến, còn M/κ là độ cứng cát tuyến; cả hai có đơn vị N·m². Các thước đo này khác với tổn hao trong phân tích cơ học động (dynamic mechanical analysis, DMA) và công vòng trễ [4], [8], [14], [18]. Do điều kiện tải và định nghĩa khác nhau, không thể gộp chúng thành một chỉ số độ cứng hoặc tiêu tán chung.')
replace_paragraph('Khi chuyển từ trạng thái',
    'Tổng hợp các mô hình sợi, lớp và cáp cho thấy quan hệ giữa trạng thái dính–trượt với độ cứng phụ thuộc đường tải và điều kiện biên [4], [5], [8]. Cần phân biệt bốn mức kết luận: ma sát hiện diện; lực pháp tuyến cục bộ đã biết; tiếp diện chi phối đã được nhận dạng; và toàn bộ vòng trễ đã được phân rã. Bằng chứng về mức đầu trong một hệ liên quan không tự động thiết lập ba mức sau cho MP1-R. Mô hình dự đoán đường cong tổng thể vì vậy cần được đối chiếu với quan sát cục bộ trước khi dùng để quy kết cơ chế.')
replace_paragraph('Bó dây MP1-R có',
    'Đối với bó dây được đề xuất, tiếp diện dây–dây và dây–ống bao là hai đường lực cản cần được xem xét khi hình học đã xác định. Các phép thử riêng cho lớp cáp và cho dây trong ống tạo tiền lệ để khảo sát từng họ tiếp diện [6], [7]. Chuyển động đầu mút, độ mềm ống bao và thay đổi cách sắp xếp dây cũng cần được kiểm soát trong MP1-R; nếu chỉ có đường cong uốn tổng thể, những khả năng này chưa được phân biệt.')
replace_paragraph('Jeon và cộng sự dùng',
    'Jeon và cộng sự đo lực kéo một dây trong vùng tiếp xúc dây–ống cao su, đồng thời khảo sát đáp ứng uốn của robot có nhiều dây làm khung chịu lực và lò xo SMA ép theo chu vi [6]. Phép kéo dây bổ sung một quan sát tiếp diện ngoài đáp ứng uốn. Tuy nhiên, lực đo là lực cản tích hợp; hệ số ma sát, lực pháp tuyến và tổn hao phụ chưa được tách độc lập. Kết quả thiết lập phép đo và điều biến lực cản trong hệ đã thử, còn việc áp dụng nó để nhận dạng dây–ống bao trong MP1-R chỉ là một tương tự phương pháp.')
replace_paragraph('Liu sử dụng các mẫu',
    'Trong khi Jeon và cộng sự xét tiếp diện dây–ống, Liu dùng mẫu cáp cắt chọn lọc để kích thích trượt giữa các lớp dây kề nhau và đo lực cản trượt tĩnh [7]. Hai phép thử tác động lên các tiếp diện khác nhau, nên có thể gợi ý cách bổ sung quan sát nhưng không phải hai phép đo cùng một đường truyền tải [6], [7]. Ở thí nghiệm của Liu, lực pháp tuyến và áp suất trung bình được suy ra từ hệ số ma sát mượn từ tài liệu cùng diện tích tiếp xúc xấp xỉ. Vì vậy, lực cản trượt được đo trực tiếp, còn ánh xạ lực pháp tuyến chưa được hiệu chuẩn độc lập.')
replace_paragraph('Hai họ bằng chứng',
    'Những tiền lệ này hỗ trợ việc bổ sung các ràng buộc tiếp diện và quan sát cục bộ để nhận dạng hệ thống (system identification) [6], [7]. Chúng không chứng minh một khoảng trống tri thức khoa học về ma sát chưa được biết. Đối với MP1-R, mục tiêu phù hợp là nhận dạng hoặc đặt cận đóng góp của các đường lực cản. Nếu dữ liệu chưa đủ để phân biệt, cần báo cáo giới hạn nhận dạng thay vì gán một nhánh tổng thể cho một tiếp diện duy nhất.')
replace_paragraph('Tương tự, một nhánh',
    'Việc liên hệ chuyển tiếp trên đường lực–độ võng với khởi phát trượt cũng cần một kiểm tra riêng. Các mô hình dính–trượt cung cấp diễn giải cơ học cho dạng nhánh [4], [5], nhưng không biến dạng nhánh thành phép đo trực tiếp chuyển động tương đối. Đối với MP1-R, quan sát trượt, biến dạng dây và chuyển động ống bao hoặc đầu mút phải hỗ trợ sự quy kết; khi thiếu chúng, chuyển tiếp chỉ được mô tả ở mức hiện tượng học.')
replace_paragraph('TiNi (NiTi) là',
    'TiNi, cũng thường được viết là NiTi, có thể thể hiện siêu đàn hồi (superelasticity) nhờ chuyển pha chịu tải và hồi phục khi dỡ tải trong miền nhiệt–cơ phù hợp. Các phép thử cáp và phần tử thành phần cho thấy đáp ứng thay đổi theo trạng thái và lịch sử [10], [11]. Hiệu ứng nhớ hình dạng (shape memory effect) được khai thác ở bộ chấp hành SMA theo một đường nhiệt–cơ khác [16]. Vì vậy, tên vật liệu TiNi không đủ để khẳng định chuyển pha đang hoạt động trong mọi miền uốn; cùng lô dây cần được đặc trưng trong miền biến dạng, nhiệt độ và lịch sử thực sự tiếp cận.')
replace_paragraph('Reedlunn, Daly',
    'Reedlunn, Daly và Shaw khảo sát gần đẳng nhiệt các cấu trúc cáp TiNi 7×7 và 1×27 cùng các mẫu thành phần, sử dụng quan sát biến dạng và nhiệt để đối chiếu đáp ứng phụ thuộc trạng thái [10], [11]. Tương quan ảnh số (digital image correlation, DIC) giúp nhận biết trường biến dạng không đồng đều bên cạnh đại lượng tải tổng thể. Các kết quả hỗ trợ tính nhạy trạng thái trong miền kéo đã thử, nhưng không xác lập ngưỡng chuyển pha hoặc tỷ phần tổn hao của bó chịu uốn dưới áp suất. Do đó, chúng định hướng kiểm soát vật liệu, không cung cấp tham số chuyển trực tiếp sang MP1-R.')
replace_paragraph('Kang và cộng sự',
    'Ở cấp mô hình, Kang và cộng sự xây dựng phần tử hữu hạn cho cáp SMA siêu đàn hồi có xét biến đổi vật liệu theo pha [12]. Khác với mô hình tiếp xúc ma sát của Vahidi và cộng sự [9], công trình này đơn giản hóa tiếp xúc và không triển khai ma sát Coulomb như một thành phần được kiểm tra độc lập. Hai cách tiếp cận cho thấy việc xét chuyển pha và việc xét lực cản tiếp diện là những lựa chọn mô hình khác nhau; sự phù hợp của một mô hình vật liệu chưa xác nhận tương tác ma sát–áp suất trong MP1-R.')
replace_paragraph('Carboni và Lacarbonara khảo sát',
    'Các bằng chứng về tiêu tán cũng cần được phân biệt theo vật liệu và phép đo. Carboni và Lacarbonara khảo sát cáp hỗn hợp NiTiNOL–thép trong bộ hấp thụ dao động, trong khi Liu và cộng sự dùng DMA và chu kỳ nhiệt cho vi sợi TiNi bện [13], [14]. Công trình thứ nhất liên hệ vòng trễ tổng hợp với ma sát và chuyển pha; công trình thứ hai cho thấy giảm chấn phụ thuộc nhiệt độ và cấu trúc bện. Sự đồng hiện của các khả năng mất mát là cơ sở để giữ nhiều cách giải thích, nhưng các phép đo này không phân tách độc lập tổn hao chuyển pha với ma sát và không tương đương công uốn chậm của bó dây.')
replace_paragraph('Từ các nguồn trên',
    'Tổng hợp thí nghiệm cáp, mô hình số và phép đo giảm chấn cho phép giữ cả nhánh vật liệu lẫn nhánh tiếp xúc trong diễn giải [9]–[15]. Tuy nhiên, bằng chứng cho sự cùng tồn tại của tổn hao không đủ để xác định tỷ phần riêng trong MP1-R. Nhiệt độ, biến dạng cục bộ, lịch sử chu kỳ và ràng buộc tiếp diện là các quan sát cần thiết nếu muốn đưa ra tuyên bố cơ chế vượt quá vòng trễ tổng hợp.')
replace_paragraph('Phase 4 đánh giá',
    'Với bằng chứng hiện có, H0a-R và H0b-R đều là các giải thích còn khả thi, còn H1-R chưa thể được phân biệt. Tính nhạy trạng thái của TiNi được ghi nhận trong những miền đã thử [10]–[14], đồng thời mô hình tiếp xúc, ma sát và giới hạn độ cứng đã có cơ sở trong các hệ liên quan [4], [5], [8], [9], [12]. Các kết quả này chưa buộc phải chọn tương tác bổ sung cho bó TiNi. Luận văn vẫn có kết quả hợp lệ nếu mô hình thông thường dự đoán tốt hoặc chuyển pha không hoạt động trong miền khảo sát.')
replace_paragraph('Ba trạng thái giả thuyết',
    'Ba cách giải thích được giữ để đánh giá trong MP1-R. H0a-R xem đáp ứng có thể được giải thích chủ yếu bởi trạng thái nhiệt–cơ của TiNi. H0b-R xem tiếp xúc, ma sát, dính–trượt, điều kiện biên và mô hình vật liệu có xét chuyển pha là đủ trong nguyên tắc, không cần luật cấu thành mới. H1-R chỉ được xem xét có điều kiện: một tương tác áp suất–tiếp xúc–vật liệu bổ sung phải được nhận dạng và cần thiết về dự đoán sau khi các cách giải thích đơn giản hơn cùng ảnh hưởng gây nhiễu đã được giới hạn.')
replace_paragraph('Huynh và cộng sự tích hợp',
    'Huynh và cộng sự tích hợp vi bơm với bộ chấp hành mềm có phần giam giữ để chuyển độ cứng [17]. Trong kiến trúc này, hút chân không phục vụ giam giữ, còn áp suất dương phục vụ uốn; hai chức năng không được đánh đồng. Kết quả cung cấp tiền lệ tích hợp nguồn áp suất và kết cấu, nhưng việc hiệu chuẩn nguồn tách rời cũng giới hạn suy luận áp suất ngay tại cụm. Nó không xác nhận miền vận hành của lò xo–piston hoặc đáp ứng bó TiNi của MP1-R.')
replace_paragraph('Các bằng chứng đã thẩm định cung cấp nhiều lớp',
    'Các mô hình hiện có khác nhau ở phần cơ học chúng giữ lại. Mô hình sợi nylon áp suất dương và mô hình dầm lớp PVC mô tả nhánh hoặc vùng dính–trượt nhưng có đường giam giữ khác nhau [4], [5]. Mô hình cáp thép cung cấp giới hạn độ cứng trong điều kiện hình học và biên đã định [8], còn các mô hình cáp SMA xét chuyển pha với mức mô tả tiếp xúc khác nhau [9], [12], [15]. Như vậy, đã có các thành phần cơ học thích hợp cho mô hình đối chứng; sự tồn tại của chúng chưa đồng nghĩa khả năng dự đoán được kiểm chứng cho bó TiNi chịu áp suất.')
replace_paragraph('Chuỗi nhân quả làm khung',
    'Chuỗi cơ học làm khung cho MP1-R gồm: áp suất chất lưu → đáp ứng màng hoặc vỏ → phản lực hướng tâm → lực pháp tuyến cục bộ → giới hạn lực ma sát → trạng thái trượt → đáp ứng uốn. Các hệ áp suất dương cung cấp bằng chứng về kích thích và đáp ứng toàn cục [3], [4]; các mô hình cáp và lớp bổ sung cơ sở cho giới hạn trượt và độ cứng [5], [8]. Việc nối các phần này thành chuỗi hoàn chỉnh của MP1-R vẫn cần những giả thiết truyền tải và quan sát riêng, như được giới hạn trong Bảng 2.3.')
replace_paragraph('Phase 5 không phê duyệt',
    'Bằng chứng đã được thẩm định không đủ để khẳng định một khoảng trống tri thức khoa học có tính khái quát cho hướng này. Phạm vi luận văn vì vậy được đặt ở hai nhu cầu đặc trưng hóa thực nghiệm, năm yêu cầu nhận dạng hệ thống, một nhu cầu xác nhận kỹ thuật nguồn và một nhu cầu xác nhận tích hợp. Việc quan hệ chưa được xác định cho một cấu hình cụ thể không đồng nghĩa quan hệ đó chưa được biết trong khoa học nói chung.')
replace_paragraph('Thứ hai, đóng góp',
    'Thứ hai, đóng góp của tiếp diện dây–dây, dây–ống bao và chuyển động ống bao hoặc đầu mút cần được đánh giá bằng quan sát độc lập. Các phép thử lực cản tiếp diện có tiền lệ ở hệ dây–ống và lớp cáp [6], [7], nhưng chưa cho phép suy ra sự chi phối của một tiếp diện trong MP1-R chỉ từ uốn tổng thể. Đây là yêu cầu nhận dạng hệ thống. Nếu quan sát chưa đủ, luận văn cần báo cáo cận hoặc giới hạn nhận dạng, thay vì chuyển giới hạn đo thành tuyên bố về một cơ chế chưa biết.')
replace_paragraph('Thứ ba, hoạt tính',
    'Thứ ba, cần đặc trưng xem chuyển pha của cùng lô TiNi có hoạt động trong miền biến dạng, nhiệt độ và lịch sử thực tế hay không. Tính nhạy trạng thái được thiết lập trong các phép thử cáp, thành phần và vi sợi đã khảo sát [10], [11], [14], còn hoạt tính trong miền uốn áp suất của MP1-R chưa được xác định. Nhu cầu này là đặc trưng hóa thực nghiệm. Phân tách chuyển pha, ma sát và tổn hao ống bao là yêu cầu nhận dạng bổ sung, chỉ khả thi khi có các quan sát độc lập đủ để phân biệt.')
replace_paragraph('Thứ tư, mô hình',
    'Thứ tư, các mô hình thông thường cung cấp đối chứng cơ học [5], [8], [9], [12], [15], nhưng khả năng dự đoán cho cấu hình MP1-R chưa được kiểm tra định lượng trên đường tải dành riêng. Nhu cầu đặc trưng hóa hiệu lực mô hình này không phải tuyên bố rằng cơ học thông thường đã bị bác bỏ. Một phần dư lặp lại chỉ liên quan đến H1-R sau khi đã kiểm soát hình học, biên, nhiệt và tiếp diện; sự cần thiết của một hạng tử tương tác mới hiện chưa được chứng minh.')
replace_paragraph('Vì vậy, tiêu đề',
    'Theo phạm vi này, tiêu đề làm việc là “Đặc trưng hóa ứng xử uốn phụ thuộc áp suất của bó dây TiNi với nguồn áp suất dẫn động bằng lò xo SMA”. Luận văn không cần H1-R được hỗ trợ: mô hình thông thường dự đoán tốt, chuyển pha không hoạt động hoặc hiệu ứng áp suất không vượt bất định đo vẫn là kết quả hợp lệ trong miền được kiểm soát.')
replace_paragraph('Đóng góp chính được giới hạn',
    'Đóng góp chính là đặc trưng hóa định lượng đáp ứng uốn phụ thuộc áp suất của bó TiNi đã định với mô hình đối chứng thông thường được khóa. Nhận dạng tiếp diện, trượt cục bộ và kiểm tra khả năng dự đoán là các đóng góp thứ cấp; xác nhận nguồn SMA–piston và tích hợp là nhiệm vụ hỗ trợ. Hạng tử tương tác bổ sung chỉ là mở rộng có điều kiện, phụ thuộc phần dư lặp lại cùng bằng chứng nhận dạng và cải thiện dự đoán trên dữ liệu giữ lại.')
anchor=next(i for i,(k,v) in enumerate(items) if k=='p' and v[0].startswith('Chuỗi cơ học làm khung'))
items.insert(anchor,('p',('Những kết quả đã được thiết lập cần được giữ trong miền kiểm chứng của chúng: giam giữ làm thay đổi đáp ứng uốn ở các hệ sợi và hạt đã thử [1], [3], [4]; trạng thái trượt tạo các giới hạn hoặc nhánh đáp ứng ở cáp thép và kết cấu nhiều lớp [5], [8]; TiNi có đáp ứng nhạy trạng thái trong các phép kéo đã khảo sát [10], [11]; và dẫn động SMA–chất lỏng đã có tiền lệ [16]. Các mô hình tiếp xúc có xét vật liệu chuyển pha cũng đã tồn tại [9], [15]. Từ các kết quả này có thể xây dựng đối chứng và phép so sánh, nhưng khả năng chuyển sang một đường phản lực mới chỉ được hỗ trợ ở mức phụ thuộc kiến trúc.',[])))

# Use consistent Vietnamese terms after their definitions.
translations = {
    'wire–wire/wire–sleeve': 'dây–dây/dây–ống bao',
    'wire–wire so với wire–sleeve': 'dây–dây so với dây–ống bao',
    'wire–wire': 'dây–dây', 'wire–sleeve': 'dây–ống bao',
    'đường điện trở': 'đường lực cản', 'tạo điện trở': 'tạo lực cản',
    'Điện trở': 'Lực cản', 'điện trở': 'lực cản',
    'sleeve/đầu mút': 'ống bao/đầu mút', 'sleeve hoặc': 'ống bao hoặc',
    'sleeve,': 'ống bao,', 'sleeve ': 'ống bao ', ' sleeve': ' ống bao',
    'packing': 'cách sắp xếp dây', 'seal': 'phớt làm kín',
    'bàng quang': 'túi bơm phồng', 'Bàng quang': 'Túi bơm phồng', 'đóng gói': 'sắp xếp phần tử',
    'giao diện': 'tiếp diện', 'trượt giao diện': 'trượt tiếp diện',
    'biến thiên theo trạng thái TiNi': 'phụ thuộc trạng thái TiNi',
    'critical weakest link được giữ nguyên từ Phase 4': 'mắt xích yếu nhất cần được kiểm tra',
}
def translate(t):
    for a,b in translations.items(): t = t.replace(a,b)
    return t

new_items=[]
for k,v in items:
    if k == 'p': new_items.append((k,(translate(v[0]),v[1])))
    elif k == 'table':
        title,heads,rows,note=v
        new_items.append((k,(translate(title),[translate(h) for h in heads],[[translate(c) for c in r] for r in rows],translate(note) if note else None)))
    else: new_items.append((k,translate(v)))
items=new_items
# Restore English equivalents at their single first definition.
replace_paragraph('Phần cuối chương kết nối',
    'Phần cuối chương kết nối chuỗi từ áp suất đến đáp ứng uốn và xác định giới hạn suy luận tại mỗi mắt xích. Trọng tâm là ánh xạ từ phản lực hướng tâm đến lực pháp tuyến tiếp xúc (contact normal force) cục bộ tại tiếp diện dây–dây (wire–wire) và dây–ống bao (wire–sleeve). Từ đó, chương dẫn đến một câu hỏi nghiên cứu chính và hai câu hỏi phụ, đồng thời giữ các cách giải thích vật liệu, tiếp xúc thông thường và tương tác bổ sung như những khả năng cần phân biệt bằng quan sát.')

for idx,(k,v) in enumerate(items):
    if k != 'table': continue
    title,heads,rows,note=v
    if title.startswith('Bảng 2.1.'):
        note='Các kết quả được diễn giải trong phạm vi từng kiến trúc gốc; riêng mô hình lớp PVC là tương tự lý thuyết–cơ học. Dòng MP1-R mô tả miền cần kiểm chứng, không phải kết quả đã có.'
    elif title.startswith('Bảng 2.2.'):
        rows=[
            ['Độ cứng F–δ','kF = dF/dδ; cát tuyến F/δ, đơn vị N/m','Đáp ứng cụm trong điều kiện gá và miền tải đã nêu','Độ cứng M–κ hoặc tiêu tán cục bộ'],
            ['Độ cứng M–κ','Tiếp tuyến dM/dκ hoặc cát tuyến M/κ, đơn vị N·m²','Ứng xử uốn của đoạn mẫu được đo','Lực pháp tuyến hoặc tiếp diện chi phối'],
            ['Tổn hao DMA','Đáp ứng động theo tần số, nhiệt độ; thí dụ tỷ số mô đun mất mát và mô đun tích trữ','Đáp ứng tiêu tán trong điều kiện phép đo','Tỷ phần ma sát trong phép uốn chậm'],
            ['Công vòng trễ','∮F dδ hoặc ∮M dθ; θ là góc quay, đơn vị rad','Công tổng hợp theo tọa độ tải–chuyển vị đã khai báo','Tỷ phần riêng của chuyển pha, trượt hoặc ống bao'],
            ['Lực cản trượt','Lực để tạo chuyển động tương đối được kiểm soát','Lực cản của tiếp diện được kích thích [6], [7]','μ và N riêng rẽ khi chưa có thông tin độc lập'],
        ]
        note='M có đơn vị N·m, κ có đơn vị m⁻¹. Diện tích ∮M dκ có đơn vị N và, khi độ cong đồng nhất, tương ứng công trên một đơn vị chiều dài; không tự nó là tổng năng lượng của mẫu. Các thước đo trong bảng phải được tách khi so sánh [4], [8], [14], [18].'
    elif title.startswith('Bảng 2.3.'):
        rows=[
            ['Áp suất chất lưu → đáp ứng màng/vỏ','Phụ thuộc mô hình','Đường áp suất phụ thuộc bao, túi và mạch; chưa có ánh xạ chung cho MP1-R [3], [4], [17].'],
            ['Màng/vỏ → phản lực hướng tâm','Được suy luận','Cần kiểm tra giả thiết truyền phản lực theo độ mềm, hình học và điều kiện biên [3]–[5].'],
            ['Phản lực hướng tâm → lực pháp tuyến dây–dây/dây–ống bao','Chưa được thiết lập','Áp suất và lực cản tổng hợp chưa xác định duy nhất lực pháp tuyến cục bộ [3]–[5], [7]. Đây là mắt xích yếu nhất.'],
            ['Lực pháp tuyến → giới hạn lực ma sát','Phụ thuộc mô hình','Luật Coulomb cung cấp một giả thiết; μ và N của tiếp diện MP1-R chưa được tách [4]–[6].'],
            ['Giới hạn lực ma sát → trạng thái trượt','Phụ thuộc mô hình','Ngưỡng và vùng dính–trượt còn phụ thuộc tải, vật liệu và biên [4], [5], [8].'],
            ['Trạng thái trượt → đáp ứng uốn','Được hỗ trợ một phần','Mô hình và phép đo tổng thể hỗ trợ sự liên hệ trong các hệ đã thử; trường trượt cục bộ của MP1-R chưa được xác nhận [4]–[6], [8].'],
        ]
        note='Bảng tổng hợp mức hỗ trợ từng liên kết; không coi chuỗi hoàn chỉnh là kết quả đã kiểm chứng trên MP1-R. Mắt xích yếu nhất là phản lực hướng tâm → lực pháp tuyến tiếp xúc cục bộ.'
    elif title.startswith('Bảng 2.4.'):
        rows=[
            ['Hình học và mặt phản lực','Tiền đề xác định cấu hình','Cần cố định trước khi diễn giải; không phải khoảng trống khoa học.'],
            ['Truyền áp suất–tiếp xúc; ảnh hưởng biên','Nhận dạng hệ thống','Xác định hoặc đặt cận truyền lực, độ mềm ống bao và chuyển động đầu mút.'],
            ['Tiếp diện và trượt cục bộ','Nhận dạng hệ thống','Nhận dạng hoặc đặt cận đường lực cản; chưa đủ quan sát thì báo cáo giới hạn nhận dạng.'],
            ['Chuyển pha cùng lô','Đặc trưng hóa thực nghiệm','Kiểm tra hoạt tính chuyển pha trong miền tiếp cận.'],
            ['Phân tách các nguồn tiêu tán','Nhận dạng hệ thống','Chỉ quy kết tỷ phần khi có quan sát độc lập đủ mạnh.'],
            ['Hiệu lực mô hình đối chứng','Đặc trưng hóa thực nghiệm','Kiểm tra mô hình khóa trên đường tải dành riêng.'],
            ['Hạng tử tương tác bổ sung','Chưa chứng minh là khoảng trống','Chỉ xem xét khi phần dư lặp lại, được nhận dạng và có giá trị dự đoán.'],
            ['Nguồn SMA–piston','Xác nhận kỹ thuật','Đặc trưng miền áp suất–thể tích và hồi phục dưới tải kết nối.'],
            ['Đáp ứng tích hợp','Xác nhận tích hợp','So sánh với nguồn ngoài dưới điều kiện ghép cặp.'],
        ]
    items[idx]=(k,(title,heads,rows,note))

# Legacy citations become work identities BEFORE assigning new numbers.
def to_tokens(t):
    t=re.sub(r'\[(\d+)\][–−-]\[(\d+)\]', lambda m: ''.join(f'{{cite:{legacy_inverse[n]}}}' for n in range(int(m[1]),int(m[2])+1)),t)
    return re.sub(r'\[(\d+)\]',lambda m:f'{{cite:{legacy_inverse[int(m[1])]}}}',t)
def map_strings(kind,value,f):
    if kind=='p':return (f(value[0]),value[1])
    if kind=='table':
        title,heads,rows,note=value
        return (f(title),[f(h) for h in heads],[[f(c) for c in r] for r in rows],f(note) if note else None)
    return f(value)
items=[(k,map_strings(k,v,to_tokens)) for k,v in items]
def texts(k,v):
    if k=='p':return [v[0]]
    if k=='table':
        title,heads,rows,note=v
        return [title,*heads,*[c for r in rows for c in r],*([note] if note else [])]
    return [v]
order={}
for k,v in items:
    for t in texts(k,v):
        for sid in re.findall(r'\{cite:([^}]+)\}',t):
            if sid not in order: order[sid]=len(order)+1
assert len(order)==18,order
def cite_render(t):
    # Merge adjacent token groups and preserve their ascending first-appearance order.
    def group(m):
        ids=re.findall(r'\{cite:([^}]+)\}',m[0]); nums=list(dict.fromkeys(order[s] for s in ids))
        nums.sort()
        chunks=[];i=0
        while i<len(nums):
            j=i
            while j+1<len(nums) and nums[j+1]==nums[j]+1:j+=1
            if j-i>=2:chunks.append(f'[{nums[i]}]–[{nums[j]}]')
            else:chunks.extend(f'[{n}]' for n in nums[i:j+1])
            i=j+1
        return ', '.join(chunks)
    return re.sub(r'\{cite:[^}]+\}(?:(?:,\s*)?\{cite:[^}]+\})*',group,t)
token_items=items
items=[(k,map_strings(k,v,cite_render)) for k,v in items]

bib={
 'S01':'L. Bai, H. Yan, J. Li, J. Shan, and P. Hou, “Detachable Soft Actuators with Tunable Stiffness Based on Wire Jamming,” Applied Sciences, vol. 12, no. 7, Art. no. 3582, 2022, doi: 10.3390/app12073582.',
 'S15':'F. Caro and M. G. Carmichael, “A Review of Mechanisms to Vary the Stiffness of Laminar Jamming Structures and Their Applications in Robotics,” Actuators, vol. 13, no. 2, Art. no. 64, 2024, doi: 10.3390/act13020064.',
 'S02':'T. Liu, H. Xia, D.-Y. Lee, A. Firouzeh, Y.-L. Park, and K.-J. Cho, “A Positive Pressure Jamming Based Variable Stiffness Structure and its Application on Wearable Robots,” IEEE Robotics and Automation Letters, vol. 6, no. 4, pp. 8078–8085, 2021, doi: 10.1109/LRA.2021.3097255.',
 'S03':'S. Zhang and J. Yao, “A variable stiffness omnidirectional chain based on positive-pressure fiber jamming,” Mechanical Sciences, vol. 17, pp. 481–493, 2026, doi: 10.5194/ms-17-481-2026.',
 'S18':'S. Zhang, J. Yao, W. Zhao, and C. Wei, “A continuum-based model for a layer jamming beam,” Mechanical Sciences, vol. 16, pp. 821–830, 2025, doi: 10.5194/ms-16-821-2025.',
 'S04':'H. Jeon, Q. N. Le, S. Jeong, S. Jang, H. Jung, H. Chang, et al., “Towards a Snake-Like Flexible Robot With Variable Stiffness Using an SMA Spring-Based Friction Change Mechanism,” IEEE Robotics and Automation Letters, vol. 7, no. 3, pp. 6582–6589, 2022, doi: 10.1109/LRA.2022.3174363.',
 'S22':'X. Liu, “Cable Vibration Considering Internal Friction,” M.S. thesis, Department of Mechanical Engineering, University of Hawaiʻi, Aug. 2004.',
 'S06':'F. Barsi, B. Carboni, and W. Lacarbonara, “A new mechanical model of short wire ropes: Theory and experimental validation,” Engineering Structures, vol. 323, Art. no. 119217, 2025, doi: 10.1016/j.engstruct.2024.119217.',
 'V-W02':'S. Vahidi, J. Arghavani, E. Choi, and A. Ostadrahimi, “Mechanical response of single and double-helix SMA wire ropes,” Mechanics of Advanced Materials and Structures, early access, Aug. 3, 2021, doi: 10.1080/15376494.2021.1955313.',
 'S05a':'B. Reedlunn, S. Daly, and J. Shaw, “Superelastic shape memory alloy cables: Part I – Isothermal tension experiments,” International Journal of Solids and Structures, 2013, doi: 10.1016/j.ijsolstr.2013.03.013.',
 'S05b':'B. Reedlunn, S. Daly, and J. Shaw, “Superelastic shape memory alloy cables: Part II – Subcomponent isothermal responses,” International Journal of Solids and Structures, 2013, doi: 10.1016/j.ijsolstr.2013.03.015.',
 'S07':'Z. Kang, Z. Wang, B. Zhou, and S. Xue, “Finite Element Method for Mechanical Behavior of Shape Memory Alloy Superelastic Cables,” Journal of Mechanical Engineering, vol. 56, no. 14, pp. 65–72, 2020, doi: 10.3901/JME.2020.14.065.',
 'S11':'B. Carboni and W. Lacarbonara, “Nonlinear Vibration Absorber with Pinched Hysteresis: Theory and Experiments,” Journal of Engineering Mechanics, Art. no. 04016023, 2016, doi: 10.1061/(ASCE)EM.1943-7889.0001072.',
 'S13':'Y. Liu, Y. Zeng, Y. Liu, Y. Xiao, T. Zhou, J. Du, et al., “High damping capacity with a wide temperature window in braided NiTi microfilaments,” Materials Letters, vol. 424, Art. no. 141544, 2026, doi: 10.1016/j.matlet.2026.141544.',
 'S12':'P. Narjabadifam, N. Fazlalipour, S. Mollaei, M. Momeni, A. S. Watandoust, M. Chavoshi, et al., “Experimental-Numerical Assessment of Mechanical Behavior of Laboratory-Made Steel and NiTi Shape Memory Alloy Wire Ropes,” Buildings, vol. 14, no. 6, Art. no. 1567, 2024, doi: 10.3390/buildings14061567.',
 'S09':'M. D. Pierce and S. A. Mascaro, “A Biologically Inspired Wet Shape Memory Alloy Actuated Robotic Pump,” IEEE/ASME Transactions on Mechatronics, vol. 18, no. 2, pp. 536–546, 2013, doi: 10.1109/TMECH.2012.2211032.',
 'S10':'H. H. Huynh, D. Han, K. Yoshida, M. De Volder, and J.-w. Kim, “Soft actuator with switchable stiffness using a micropump-activated jamming system,” Sensors and Actuators A: Physical, vol. 338, Art. no. 113449, 2022, doi: 10.1016/j.sna.2022.113449.',
 'S17':'B. Lin, S. Song, and J. Wang, “Variable stiffness methods of flexible robots for minimally invasive surgery: A review,” Biomimetic Intelligence and Robotics, vol. 4, Art. no. 100168, 2024, doi: 10.1016/j.birob.2024.100168.',
}
assert set(bib)==set(order)

# Detailed evidence map saved before building the revised document.
section_specs={
 '2.1':('Position the bounded review and architecture','confinement / structural material / pressure source','external study domains versus proposed architecture',['RP017','RP019'],['GAP02','GAP08'],['CONTRIB01']),
 '2.2':('Compare confinement architectures and pressure paths','vacuum wire, positive-pressure granular/fiber, vacuum PVC layer','material, pressure direction, reaction surfaces, loading metric',['RP001','RP002','RP003','RP008','RP016','RP021'],['GAP01','GAP02'],['CONTRIB01']),
 '2.3':('Establish slip/friction comparator and metric discipline','friction bound, slip states, rope stiffness limits','global response versus local forces; F–δ versus M–κ versus losses',['RP005','RP007','RP008','RP011','RP015','RP017'],['GAP02','GAP08'],['CONTRIB01','CONTRIB02']),
 '2.4':('Bound interface and slip identifiability','internal cable-layer tests versus wire–rubber pull tests','target interface, intervention, measured resistance versus inferred normal load',['RP009','RP013','RP018'],['GAP04','GAP05'],['CONTRIB03','CONTRIB04']),
 '2.5':('Relate structural TiNi response to accessed domain','cable/component tension, DMA, transformation-aware numerics','experimental NiTi evidence versus steel tests and numerical SMA',['RP006','RP007','RP012'],['GAP06','GAP07'],['CONTRIB01','CONTRIB02']),
 '2.6':('Preserve competing explanations and conditional H1','material state versus conventional material/contact baseline','coexisting aggregate losses versus separately identified mechanisms',['RP012','RP019'],['GAP07','GAP09'],['CONTRIB01','CONTRIB02','CONTRIB07']),
 '2.7':('Separate actuator SMA from structural TiNi','SMA-to-fluid and SMA friction-control precedents','pump architecture, vacuum locking versus positive bending, thermal roles',['RP004','RP010','RP020'],['GAP10'],['CONTRIB05']),
 '2.8':('Explain Bench A/B as project design reasoning','separate structural response from source dynamics','matched specimen pressure, temperature and history versus unmatched comparison',['RP014','RP020'],['GAP10','GAP11'],['CONTRIB05','CONTRIB06']),
 '2.9':('Motivate conventional predictive comparator','branch/continuum/rope/material formulations','different model assumptions; calibration versus reserved validation',['RP005','RP007','RP008','RP011','RP019'],['GAP08','GAP09'],['CONTRIB01','CONTRIB02']),
 '2.10':('Synthesize causal-chain support and limits','linkwise causal assessment','architecture evidence versus model-dependent transmission',['RP011','RP016','RP017','RP018'],['GAP02','GAP04'],['CONTRIB01','CONTRIB03']),
 '2.11':('State approved needs without literature-wide gap claim','2 empirical needs, 5 identification requirements, source/integration needs','scope prerequisite versus scientific gap; unestablished coupling versus required new law',['RP017','RP018','RP019','RP020'],[f'GAP{i:02}' for i in range(1,12)],['CONTRIB01','CONTRIB02','CONTRIB03','CONTRIB05','CONTRIB06']),
 '2.12':('Connect synthesis to the approved RQ and contribution boundary','one primary and two secondary questions','core thesis success versus optional integration/coupling',['RP017','RP018','RP019','RP020'],['GAP02','GAP04','GAP06','GAP08','GAP09'],['CONTRIB01','CONTRIB02','CONTRIB03','CONTRIB07']),
}
canonical_strength={**{f'RP{i:03}':'MODERATE' for i in range(1,22)},'RP001':'STRONG','RP002':'STRONG','RP012':'LIMITED',**{f'RP{i:03}':'ANALOGUE_ONLY' for i in [13,14,15,16]},**{f'RP{i:03}':'NOT_ESTABLISHED' for i in [17,18,19,20]}}
section=''; evidence_map=[]
for k,v in token_items:
    if k=='h2':
        section=v.split('.')[0]+'.'+v.split('.')[1]
        purpose,claims,compare,rps,gaps,contrib=section_specs[section]
        evidence_map.append({'subsection':section,'title':v,'purpose':purpose,'claims':claims,'sources':[], 'evidence_strength':{rp:canonical_strength[rp] for rp in rps},'comparison':compare,'contrast':compare,'limitations':'No direct transfer of source parameters or mechanisms to MP1-R; retain each source-specific restriction.','mp1r_connection':purpose,'prohibited_overclaim':p5['prohibited_claims'],'phase4_propositions':rps,'phase5_gaps':gaps,'contributions':contrib})
    if evidence_map:
        for t in texts(k,v):
            evidence_map[-1]['sources'].extend(re.findall(r'\{cite:([^}]+)\}',t))
for m in evidence_map:m['sources']=list(dict.fromkeys(m['sources']))
(RUN/'chapter2_evidence_map.json').write_text(json.dumps(evidence_map,ensure_ascii=False,indent=2))
(RUN/'chapter2_content_inventory.json').write_text(json.dumps({'items':items,'reference_order':order,'bibliography':bib},ensure_ascii=False,indent=2))

doc=Document();sec=doc.sections[0]
sec.page_width=Mm(210);sec.page_height=Mm(297)
sec.top_margin=Cm(3);sec.bottom_margin=Cm(3);sec.left_margin=Cm(3.5);sec.right_margin=Cm(2)
sec.header_distance=Cm(1.3);sec.footer_distance=Cm(1.3)
def font(run,size=13,bold=None,italic=None):
    run.font.name='Times New Roman';run.font.size=Pt(size)
    run.font.color.rgb=RGBColor(0,0,0)
    if bold is not None:run.bold=bold
    if italic is not None:run.italic=italic
    pr=run._r.get_or_add_rPr()
    for key in ['ascii','hAnsi','eastAsia','cs']:pr.rFonts.set(qn('w:'+key),'Times New Roman')
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'vi-VN');pr.append(lang)
def para(p,align=WD_ALIGN_PARAGRAPH.JUSTIFY,first=1,line=1.5):
    p.alignment=align;pf=p.paragraph_format;pf.first_line_indent=Cm(first)
    pf.line_spacing=line;pf.space_before=Pt(0);pf.space_after=Pt(0)
for name,size,align in [('Normal',13,WD_ALIGN_PARAGRAPH.JUSTIFY),('Heading 1',16,WD_ALIGN_PARAGRAPH.CENTER),('Heading 2',13,WD_ALIGN_PARAGRAPH.LEFT),('Heading 3',13,WD_ALIGN_PARAGRAPH.LEFT)]:
    st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size)
    st.font.color.rgb=RGBColor(0,0,0)
    rpr=st.element.get_or_add_rPr()
    for key in ['ascii','hAnsi','eastAsia','cs']:rpr.rFonts.set(qn('w:'+key),'Times New Roman')
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'vi-VN');rpr.append(lang)
    pf=st.paragraph_format;pf.alignment=align;pf.line_spacing=1.5;pf.first_line_indent=Cm(1 if name=='Normal' else 0);pf.space_after=Pt(0 if name=='Normal' else 6)
    if name!='Normal':st.font.bold=True;pf.keep_with_next=True;pf.space_before=Pt(12 if name!='Heading 1' else 0)
header=sec.header.paragraphs[0];para(header,WD_ALIGN_PARAGRAPH.CENTER,0)
r=header.add_run();font(r,13)
for typ in ['begin','separate','end']:
    fld=OxmlElement('w:fldChar');fld.set(qn('w:fldCharType'),typ);r._r.append(fld)
    if typ=='begin':
        instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=' PAGE ';r._r.append(instr)
    elif typ=='separate':t=OxmlElement('w:t');t.text='1';r._r.append(t)
pg=OxmlElement('w:pgNumType');pg.set(qn('w:start'),'1');sec._sectPr.append(pg)
doc.settings.element.append(OxmlElement('w:updateFields'))
doc.settings.element[-1].set(qn('w:val'),'true')
doc.core_properties.title='Chương 2. Tổng quan nghiên cứu và cơ sở khoa học — MP1-R'
doc.core_properties.subject=p5['working_title_audit']['conservative_title']
doc.core_properties.language='vi-VN'
doc.core_properties.author=''
doc.core_properties.last_modified_by=''
doc.core_properties.comments='Evidence-controlled standalone Chapter 2 draft. Page numbering starts at 1.'

for k,v in items:
    if k in ['h1','h2','h3']:
        p=doc.add_paragraph(v,style={'h1':'Heading 1','h2':'Heading 2','h3':'Heading 3'}[k])
        for r in p.runs:font(r,16 if k=='h1' else 13,True)
    elif k=='p':
        p=doc.add_paragraph();para(p);font(p.add_run(v[0]))
    elif k=='eq':
        p=doc.add_paragraph();para(p,WD_ALIGN_PARAGRAPH.CENTER,0)
        math=OxmlElement('m:oMath')
        def mr(txt):
            r=OxmlElement('m:r');pr=OxmlElement('w:rPr');rf=OxmlElement('w:rFonts')
            rf.set(qn('w:ascii'),'Times New Roman');rf.set(qn('w:hAnsi'),'Times New Roman');pr.append(rf)
            sz=OxmlElement('w:sz');sz.set(qn('w:val'),'26');pr.append(sz)
            color=OxmlElement('w:color');color.set(qn('w:val'),'000000');pr.append(color);r.append(pr)
            t=OxmlElement('m:t');t.text=txt;r.append(t);return r
        def sub(base,index):
            s=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(mr(base));s.append(e)
            ix=OxmlElement('m:sub');ix.append(mr(index));s.append(ix);return s
        delimiter=OxmlElement('m:d');pr=OxmlElement('m:dPr')
        for tag in ['m:begChr','m:endChr']:
            c=OxmlElement(tag);c.set(qn('m:val'),'|');pr.append(c)
        delimiter.append(pr);e=OxmlElement('m:e');e.append(sub('F','t,i'));delimiter.append(e)
        math.append(delimiter);math.append(mr(' ≤ '));math.append(sub('μ','i'));math.append(mr(' '));math.append(sub('N','i'))
        p._p.append(math)
        font(p.add_run('     (2.1)'))
    elif k=='table':
        title,heads,rows,note=v
        p=doc.add_paragraph();para(p,WD_ALIGN_PARAGRAPH.CENTER,0,1.2);p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(4)
        font(p.add_run(title),12,True)
        table=doc.add_table(rows=1,cols=len(heads));table.style='Table Grid';table.alignment=WD_TABLE_ALIGNMENT.CENTER
        for j,row in enumerate([heads,*rows]):
            rr=table.rows[0] if j==0 else table.add_row()
            trpr=rr._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
            if j==0:
                rep=OxmlElement('w:tblHeader');rep.set(qn('w:val'),'true');trpr.append(rep)
            for c,txt in zip(rr.cells,row):
                c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
                p=c.paragraphs[0];para(p,WD_ALIGN_PARAGRAPH.LEFT,0,1)
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                font(p.add_run(txt),11,j==0)
        if note:
            p=doc.add_paragraph();para(p,WD_ALIGN_PARAGRAPH.JUSTIFY,0,1.2)
            font(p.add_run(note),11,italic=True);p.paragraph_format.space_after=Pt(6)
    elif k=='note':
        p=doc.add_paragraph();para(p,WD_ALIGN_PARAGRAPH.JUSTIFY,0,1.2);font(p.add_run(v),11,italic=True)

# Bibliography heading already exists as the last inventory item.
doc.paragraphs[-1].paragraph_format.page_break_before=True
for sid,n in order.items():
    p=doc.add_paragraph();para(p,WD_ALIGN_PARAGRAPH.JUSTIFY,0)
    p.paragraph_format.left_indent=Cm(0.8);p.paragraph_format.first_line_indent=Cm(-0.8)
    font(p.add_run(f'[{n}] {bib[sid]}'))
doc.save(ROOT/DOCX)

# Trace each substantive prose paragraph, equation, table row and table note.
source_rps={'S01':['RP001'],'S02':['RP001','RP002','RP016'],'S03':['RP001','RP002','RP003','RP011','RP016'],'S04':['RP004','RP009','RP011','RP013'],'S22':['RP009','RP013'],'S06':['RP005','RP011'],'S05a':['RP006','RP012'],'S05b':['RP006','RP012'],'S07':['RP019'],'S11':['RP012'],'S12':['RP007'],'S13':['RP012','RP015'],'S09':['RP010'],'S10':['RP014'],'S15':['RP021'],'S17':['RP015'],'S18':['RP008','RP011','RP016'],'V-W02':['RP007','RP011']}
canonical=lambda s:'S05' if s in ['S05a','S05b'] else s
traces=[];section='';para_index=0
for (k,token),(kk,rendered) in zip(token_items,items):
    if k=='h2':section=token.split('.')[0]+'.'+token.split('.')[1];para_index=0
    if k not in ['p','eq','table','note']:continue
    if k=='p':units=[('paragraph',token[0],rendered[0])];para_index+=1
    elif k=='table':
        units=[(f'table_row_{i+1}',' | '.join(row),' | '.join(rendered[2][i])) for i,row in enumerate(token[2])]
        if token[3]:units.append(('table_note',token[3],rendered[3]))
    else:units=[(k,token,rendered)]
    spec=section_specs[section]
    for unit,txt,final in units:
        sids=list(dict.fromkeys(re.findall(r'\{cite:([^}]+)\}',txt)))
        inherited=False
        if not sids and k=='eq':
            sids=['S03','S18'];inherited=True
        if not sids and k=='table' and token[0].startswith('Bảng 2.2.'):
            sids=['S03','S06','S13','S17'];inherited=True
        candidates={rp for s in sids for rp in source_rps[s]}
        rps=[rp for rp in spec[3] if rp in candidates]
        project=not sids
        if project:rps=spec[3]
        # Explicit bounded unresolved statements are linked to their unresolved RP.
        if any(w in final for w in ['pháp tuyến','truyền áp suất','áp suất–tiếp xúc']) and section in ['2.3','2.10','2.11']:rps=list(dict.fromkeys([*rps,'RP017']))
        if any(w in final for w in ['chi phối','nhận dạng','cục bộ']) and section in ['2.4','2.11']:rps=list(dict.fromkeys([*rps,'RP018']))
        if any(w in final for w in ['đối chứng','H1-R','luật cấu thành','mô hình thông thường']) and section in ['2.6','2.9','2.11','2.12']:rps=list(dict.fromkeys([*rps,'RP019']))
        if final.startswith('Những kết quả đã được thiết lập cần'):
            rps=['RP001','RP002','RP005','RP006','RP007','RP008','RP010','RP016']
        gaps=spec[4];contrib=spec[5]
        if section=='2.11' and k=='p':
            gaps={1:spec[4],2:['GAP01','GAP02','GAP03'],3:['GAP04','GAP05'],4:['GAP06','GAP07'],5:['GAP08','GAP09'],6:['GAP10','GAP11']}[para_index]
        if k=='table' and token[0].startswith('Bảng 2.4.') and unit.startswith('table_row'):
            gaps={1:['GAP01'],2:['GAP02','GAP03'],3:['GAP04','GAP05'],4:['GAP06'],5:['GAP07'],6:['GAP08'],7:['GAP09'],8:['GAP10'],9:['GAP11']}[int(unit.rsplit('_',1)[1])]
        if section=='2.12' and k=='p':
            gaps={1:['GAP02','GAP03','GAP08'],2:['GAP04','GAP05','GAP06','GAP07'],3:['GAP02','GAP04','GAP06','GAP08','GAP09','GAP10','GAP11'],4:[]}[para_index]
        # Unresolved RP entries denote claim boundaries, not strength for a negative scientific law.
        strength={rp:canonical_strength[rp] for rp in rps}
        cls='inference' if project else ('synthesis' if len(sids)>1 else 'direct')
        if section in ['2.8']:cls='inference'
        if any(w in final for w in ['tương tự','tương đồng','tham chiếu']):cls='analogy'
        traces.append({'claim_id':f'CH2C{len(traces)+1:03}','chapter_subsection':section,'location':unit+(f'_{para_index}' if unit=='paragraph' else ''),'claim_text':final,'source_ids':list(dict.fromkeys(canonical(s) for s in sids)),'source_work_ids':sids,'ieee_reference_numbers':[order[s] for s in sids],'citation_inherited_from_lead_or_table_note':inherited,'evidence_status':strength or {'policy':'PHASE5_APPROVED_SCOPE_OR_MEASUREMENT_DEFINITION'},'classification':cls,'phase4_proposition_mapping':rps,'phase5_gap_mapping':gaps,'phase5_contribution_mapping':contrib,'wording_boundary':'Source findings remain scoped to their own architecture/domain; project need is not a scientific absence claim. No direct MP1-R mechanism/result is asserted.','authority_basis':P5PATH if project else 'Phase 3 appraisal and canonical Phase 4 propositions; original source for verified bibliography.','audit_status':'DRAFT_TIME_CHECKED_PHASE7_PENDING'})
trace={'schema_version':'1.1','artifact_type':'phase6_chapter2_claim_traceability','date':DATE,'workstream':'MP1-R','docx_path':DOCX,'primary_input_manifest':manifest,'evidence_map_path':str((RUN/'chapter2_evidence_map.json').relative_to(ROOT)),'claim_count':len(traces),'claims':traces,'reference_register':[{'number':n,'source_id':canonical(s),'source_work_id':s,'bibliographic_entry':bib[s],'restriction':{'S18':'vacuum PVC layer-jamming mechanistic analogue only','S15':'secondary laminar review context; excludes fiber/particle coverage','S17':'metric-definition/context only','V-W02':'published numerical cable paper; 2021 online version cited','S07':'transformation-aware FE; frictionless/simplified contact','S12':'steel experiments; SMA numerical only'}.get(s,'Retain Phase 3 source-specific restrictions')} for s,n in order.items()],'phase7_audit_executed':False}
(ROOT/TRACE).write_text(json.dumps(trace,ensure_ascii=False,indent=2))

reopened=['S01','S02','S03','S04','S05a','S06','S07','S09','S10','S11','S12','S13','S17','V-W02']
entries=[]
for sid in reopened:
    path=ROOT/f'outputs/reports/MP1_R_PHASE_3_2026-10-02/{"W02_VAHIDI" if sid=="V-W02" else sid}_fulltext.txt'
    entries.append({'source_id':canonical(sid),'source_work_id':sid,'accessed_artifact':str(path.relative_to(ROOT)),'artifact_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'access_type':'archived original-publication text, targeted first-page/header read','reason':'Verify actual authors and bibliographic identity before IEEE finalization; do not infer author initials from a surname-only authority summary.','claim_checked':'Author sequence, title, journal/year and available volume/page/article fields; no re-adjudication of scientific propositions.','section_page':'Publisher cover/first article page; printed page headers additionally checked for S04 (6582–6589), S07 (65–72), S09 (536–546).','result':'Corrected bibliographic record in reference register; scientific boundary unchanged.','special_result':{'S17':'Journal corrected to Biomimetic Intelligence and Robotics, vol.4, article100168; metric-context role preserved.','V-W02':'Publisher cover states published online 3 Aug 2021; cite verified early-access version, not unsupported issue-year/pages.','S03':'First-page header verifies Mechanical Sciences 17, 481–493 (2026).','S05a':'Names verified as Benjamin Reedlunn, Samantha Daly, John Shaw; final issue pages not inferred.'}.get(sid,None)})
log={'schema_version':'1.1','artifact_type':'phase6_source_reopen_log','date':DATE,'workstream':'MP1-R','total_reopened_sources':14,'total_reopened_works':14,'canonical_source_ids':list(dict.fromkeys(canonical(s) for s in reopened)),'reopen_log':entries,'not_reopened_original_works':['S05b','S15','S18','S22'],'other_metadata_basis':'Existing Phase 3 bibliography and source provenance/verified evidence records; no new-source discovery.','broad_literature_searches':0,'targeted_external_literature_searches':0,'new_sources_added':0,'scientific_reconciliation_repeated':False}
(ROOT/REOPEN).write_text(json.dumps(log,ensure_ascii=False,indent=2))

bodytexts=[]
for k,v in items:
    if k=='h1' and v=='TÀI LIỆU THAM KHẢO':break
    bodytexts.extend(texts(k,v))
body='\n'.join(bodytexts)
metrics={'chapter_sections':12,'unique_references':len(order),'canonical_sources_cited':len(set(canonical(s) for s in order)),'in_text_citations':len(re.findall(r'\[\d+\](?:[–−-]\[\d+\])?',body)),'citation_number_tokens':len(re.findall(r'\[\d+\]',body)),'body_words':len(body.split()),'body_characters':len(body),'tables':4,'equations':1,'claim_count':len(traces),'original_papers_reopened':14,'broad_literature_searches':0,'counting_rule':'Citation appearances count one single [n] or one range [n]–[m] as one; separate [n], [m] count two. Includes all body/table captions, cells and notes; excludes bibliography.'}
(ROOT/'outputs/plans/MP1_R_PHASE_6_DRAFT_METRICS.json').write_text(json.dumps(metrics,ensure_ascii=False,indent=2))
(RUN/'chapter2_readable_draft.md').write_text('\n\n'.join(('## '+v if k in ['h1','h2'] else v[0] if k=='p' else '\n'.join([v[0],' | '.join(v[1]),*[' | '.join(r) for r in v[2]],v[3] or '']) if k=='table' else v) for k,v in items)+'\n\n'+'\n\n'.join(f'[{n}] {bib[s]}' for s,n in order.items()))
print(json.dumps({'docx':DOCX,'reference_order':order,'metrics':metrics},ensure_ascii=False,indent=2))
