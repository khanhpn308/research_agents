from pathlib import Path
import json, re, zipfile, hashlib
from lxml import etree
from docx import Document
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[3]; RUN=Path(__file__).parent
docx=ROOT/'docs/project/MP1_R_CHAPTER_2_EVIDENCE_CONTROLLED_DRAFT_2026-10-02.docx'
inventory=json.loads((RUN/'chapter2_content_inventory.json').read_text())
trace=json.loads((ROOT/'outputs/plans/MP1_R_PHASE_6_CHAPTER_2_CLAIM_TRACEABILITY_2026-10-02.json').read_text())
reopen=json.loads((ROOT/'outputs/plans/MP1_R_PHASE_6_SOURCE_REOPEN_LOG_2026-10-02.json').read_text())
checks={}
def check(k,condition,detail=None):
    checks[k]={'result':'PASS' if condition else 'FAIL','detail':detail}
    assert condition,(k,detail)

with zipfile.ZipFile(docx) as z:
    check('valid_ooxml_package',z.testzip() is None)
    for name in [x for x in z.namelist() if x.endswith('.xml')]:etree.fromstring(z.read(name))
    xml=etree.fromstring(z.read('word/document.xml'))
    header=etree.fromstring(z.read('word/header1.xml'))
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
d=Document(docx);s=d.sections[0]
check('docx_opens_python_docx',len(d.paragraphs)>0)
check('A4',abs(s.page_width.mm-210)<0.05 and abs(s.page_height.mm-297)<0.05)
actual={k:getattr(s,k+'_margin').cm for k in ['top','bottom','left','right']}
check('margins',all(abs(actual[k]-v)<0.003 for k,v in {'top':3,'bottom':3,'left':3.5,'right':2}.items()),actual)
normal=d.styles['Normal']
check('body_font_Times_New_Roman',normal.font.name=='Times New Roman')
check('body_font_13pt',normal.font.size.pt==13)
check('body_line_spacing_1_5',normal.paragraph_format.line_spacing==1.5)
check('body_justified',normal.paragraph_format.alignment==WD_ALIGN_PARAGRAPH.JUSTIFY)
check('first_line_indent_1cm',abs(normal.paragraph_format.first_line_indent.cm-1)<0.003)
check('Vietnamese_language',normal.element.xpath('.//w:lang')[0].get(qn('w:val'))=='vi-VN')
check('page_number_top_center',any('PAGE' in x for x in header.xpath('//w:instrText/text()',namespaces=ns)) and header.xpath('//w:jc/@w:val',namespaces=ns)==['center'])
check('page_number_start_1',xml.xpath('//w:pgNumType/@w:start',namespaces=ns)==['1'])
h2=[p.text for p in d.paragraphs if p.style.name=='Heading 2']
check('heading_hierarchy',len(h2)==12 and all(t.startswith(f'2.{i}.') for i,t in enumerate(h2,1)),h2)
check('chapter_heading',d.paragraphs[0].style.name=='Heading 1' and d.paragraphs[0].text.startswith('CHƯƠNG 2.'))
check('editable_tables',len(d.tables)==4 and len(xml.xpath('//w:tbl',namespaces=ns))==4)
check('editable_equation',len(xml.xpath('//m:oMath',namespaces=ns))==1)
check('no_rasterized_content',len(xml.xpath('//w:drawing|//w:pict',namespaces=ns))==0)
check('references_heading',any(p.text=='TÀI LIỆU THAM KHẢO' for p in d.paragraphs))

# Traverse the REAL DOCX order, including tables (not d.paragraphs alone).
body=[]; in_refs=False; refs=[]
for child in xml.find(qn('w:body')):
    txt=''.join(child.xpath('.//w:t/text()',namespaces=ns))
    if txt=='TÀI LIỆU THAM KHẢO':in_refs=True;continue
    if in_refs:
        if txt:refs.append(txt)
    elif child.tag==qn('w:tbl'):
        for cell in child.xpath('./w:tr/w:tc',namespaces=ns):body.append(''.join(cell.xpath('.//w:t/text()',namespaces=ns)))
    else:body.append(txt)
text='\n'.join(body);sequence=[];seen=set()
for match in re.finditer(r'\[(\d+)\](?:[–−-]\[(\d+)\])?',text):
    nums=range(int(match[1]),int(match[2])+1) if match[2] else [int(match[1])]
    for n in nums:
        if n not in seen:seen.add(n);sequence.append(n)
check('IEEE_first_appearance_order',sequence==list(range(1,19)),sequence)
check('reference_number_match',[int(re.match(r'\[(\d+)\]',t)[1]) for t in refs]==sequence)
check('no_orphan_references_or_citations',len(refs)==18 and seen==set(range(1,19)))
check('no_author_year_citations',not re.search(r'\([^)]*\b(?:19|20)\d{2}\b[^)]*\)|\[(?:[A-Z][a-z]+),',text))
check('no_internal_registry_ids',not re.search(r'\b(?:RP\d{3}|GAP\d{2}|CONTRIB\d{2}|Phase [3-5]|NOT_ESTABLISHED|PLAUSIBLE)\b',text))
check('no_excluded_source_citations',not re.search(r'\bS(?:19|20|21)\b',text) and not any(x['source_id'] in ['S19','S20','S21'] for x in trace['reference_register']))
check('no_unconditional_novelty_language',not re.search(r'(nghiên cứu đầu tiên|lần đầu tiên|chưa có nghiên cứu nào|không có nghiên cứu nào|hoàn toàn mới)',text,re.I))
check('no_literature_discovery',reopen['broad_literature_searches']==0 and reopen['new_sources_added']==0)
allowed=set('S01 S02 S03 S04 S05 S06 S07 S09 S10 S11 S12 S13 S15 S17 S18 S22 V-W02'.split())
check('source_IDs_resolve',all(set(c['source_ids'])<=allowed for c in trace['claims']))
check('claim_register_complete',trace['claim_count']==len(trace['claims'])==84)
register={r['source_work_id']:r['number'] for r in trace['reference_register']}
check('claim_reference_mapping',all(c['ieee_reference_numbers']==[register[s] for s in c['source_work_ids']] for c in trace['claims']))
doi=[re.search(r'doi:\s*(\S+)\.$',t)[1] for t in refs if 'doi:' in t]
check('no_duplicate_bibliographic_DOI',len(doi)==len(set(doi))==17)
s18=next(r for r in trace['reference_register'] if r['source_id']=='S18')
check('S18_correct_identity','A continuum-based model for a layer jamming beam' in s18['bibliographic_entry'] and '10.5194/ms-16-821-2025' in s18['bibliographic_entry'])
check('source_reopen_log_count',reopen['total_reopened_sources']==len(reopen['reopen_log'])==14)
manifest=json.loads((RUN/'input_manifest.json').read_text())
check('Phase3_to_5_inputs_immutable',all(hashlib.sha256((ROOT/x['path']).read_bytes()).hexdigest()==x['sha256'] for x in manifest))
for sid,expected in [('S15','context'),('S17','context'),('S18','analogue')]:
    check(f'{sid}_restriction_recorded',expected in next(r for r in trace['reference_register'] if r['source_id']==sid)['restriction'])

# Opening in a full office renderer is stronger than ZIP/parse validation alone.
pdf=RUN/'render/MP1_R_CHAPTER_2_EVIDENCE_CONTROLLED_DRAFT_2026-10-02.pdf'
check('office_renderer_opened',pdf.exists() and pdf.read_bytes().startswith(b'%PDF'))
report={'schema_version':'1.0','artifact_type':'phase6_draft_time_validation','docx_path':str(docx.relative_to(ROOT)),'docx_sha256':hashlib.sha256(docx.read_bytes()).hexdigest(),'checks':checks,'automated_check_count':len(checks),'automated_result':'PASS','visual_validation':'PENDING','phase7_audit_executed':False,'limitations':['This is draft-time validation, not the independent Phase 7 claim/citation audit.','Microsoft Word was not available; DOCX opened using python-docx and LibreOffice 24.2.7.2.','Physical margins validated with Word twip-rounding tolerance.']}
(RUN/'formatting_and_content_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps({'automated_result':'PASS','checks':len(checks),'first_appearance':sequence,'margins':actual},ensure_ascii=False))
