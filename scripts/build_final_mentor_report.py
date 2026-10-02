# -*- coding: utf-8 -*-
"""
Script to build the exact, concise, citation-grounded mentor report:
1. Markdown: docs/reports/D1_MENTOR_RESEARCH_NARROWING_REPORT.md
2. DOCX: docs/reports/D1_MENTOR_RESEARCH_NARROWING_REPORT.docx

Conforms strictly to:
- 8 Exact Required Sections
- Target: 4-6 pages main content + References (Max 7 pages main content)
- Font: Times New Roman throughout
- Body: 13 pt, Line spacing 1.3, Space after 3 pt, Justified
- Heading 1: 14 pt Bold, Line spacing 1.3, Space before 8 pt, Space after 2 pt
- Heading 2: 13 pt Bold, Line spacing 1.3, Space before 5 pt, Space after 2 pt
- Table font: 12 pt Times New Roman (headers 12 pt bold), cell padding, cantSplit, tblHeader, light gray shading
- Header: Right-aligned 8.5 pt italic gray
- Footer: Center-aligned page numbers (Page X of Y)
- Style: Short, clear, technical, factual, evidence-first, zero conversational fluff
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from report_helpers import (
    set_cell_margins, set_cell_shading, set_table_borders,
    make_row_header, prevent_row_split, add_page_number_to_run, add_total_pages_to_run
)

def build_report():
    doc = Document()
    
    # Page Setup: A4, Margins: Top 2.0cm, Bottom 2.0cm, Left 2.5cm, Right 2.0cm
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)
    
    # Styles Configuration
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(13)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.3
    normal_style.paragraph_format.space_after = Pt(3)
    normal_style.paragraph_format.space_before = Pt(0)
    normal_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    # Header & Footer
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("Báo cáo Thu hẹp Phạm vi Nghiên cứu D1 — Trình Cán bộ hướng dẫn (26/09/2026)")
    hrun.font.name = 'Times New Roman'
    hrun.font.size = Pt(8.5)
    hrun.font.italic = True
    hrun.font.color.rgb = RGBColor(120, 120, 120)
    
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun1 = fp.add_run("Trang ")
    frun1.font.name = 'Times New Roman'
    frun1.font.size = Pt(9.0)
    frun1.font.color.rgb = RGBColor(100, 100, 100)
    add_page_number_to_run(frun1)
    frun2 = fp.add_run(" / ")
    frun2.font.name = 'Times New Roman'
    frun2.font.size = Pt(9.0)
    frun2.font.color.rgb = RGBColor(100, 100, 100)
    add_total_pages_to_run(frun2)
    
    def add_h1(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0, 32, 96)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(40, 40, 40)
        return p

    def add_p(text, bold_prefix="", italic_prefix=""):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(3)
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = 'Times New Roman'
            rb.font.size = Pt(13)
            rb.font.bold = True
        if italic_prefix:
            ri = p.add_run(italic_prefix)
            ri.font.name = 'Times New Roman'
            ri.font.size = Pt(13)
            ri.font.italic = True
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13)
        return p

    def format_table(tbl, col_widths, font_size=12.0):
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        for j, w in enumerate(col_widths):
            tbl.columns[j].width = w
        set_table_borders(tbl, color="B8B8B8", sz="4", val="single")
        for i, row in enumerate(tbl.rows):
            prevent_row_split(row)
            if i == 0:
                make_row_header(row)
            for j, cell in enumerate(row.cells):
                cell.width = col_widths[j]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=30, bottom=30, left=45, right=45)
                if i == 0:
                    set_cell_shading(cell, "EAECEE")
                else:
                    if i % 2 == 1:
                        set_cell_shading(cell, "FFFFFF")
                    else:
                        set_cell_shading(cell, "F8F9FA")
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.15
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    for run in p.runs:
                        run.font.name = 'Times New Roman'
                        run.font.size = Pt(font_size)
                        if i == 0:
                            run.font.bold = True

    # -------------------------------------------------------------
    # TITLE & METADATA
    # -------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(1)
    p_title.paragraph_format.line_spacing = 1.15
    run_t = p_title.add_run("BÁO CÁO KHOA HỌC: TIẾN TRÌNH THU HẸP PHẠM VI NGHIÊN CỨU\nTỪ CHỦ ĐỀ TỔNG QUAN ĐẾN CÂU HỎI NGHIÊN CỨU D1")
    run_t.font.name = 'Times New Roman'
    run_t.font.size = Pt(14)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(0, 32, 96)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(3)
    run_s = p_sub.add_run("(GIỚI HẠN HIỆU LỰC MÔ HÌNH CONTINUUM DẦM LAYER JAMMING)")
    run_s.font.name = 'Times New Roman'
    run_s.font.size = Pt(11.5)
    run_s.font.bold = True
    run_s.font.color.rgb = RGBColor(70, 70, 70)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_before = Pt(0)
    p_meta.paragraph_format.space_after = Pt(6)
    p_meta.paragraph_format.line_spacing = 1.1
    run_m = p_meta.add_run("Báo cáo tiến trình rà soát y văn trình Cán bộ hướng dẫn (Mentor Progress Report)\n"
                           "Chuyên ngành: Cơ kỹ thuật / Robot mềm | Thời điểm: 26/09/2026\n"
                           "Phương pháp luận: Falsification-First & Evidence-Based Narrowing")
    run_m.font.name = 'Times New Roman'
    run_m.font.size = Pt(9.0)
    run_m.font.italic = True
    run_m.font.color.rgb = RGBColor(90, 90, 90)

    # -------------------------------------------------------------
    # SECTION 1 — MỤC TIÊU BAN ĐẦU
    # -------------------------------------------------------------
    add_h1("SECTION 1 — MỤC TIÊU BAN ĐẦU")
    add_p(
        "Robot mềm (soft robotics) và bài toán điều biến độ cứng thích ứng (variable stiffness).",
        bold_prefix="Chủ đề ban đầu (Initial Topic): "
    )
    add_p(
        "Khái niệm 'nghiên cứu dầm layer jamming' là một vùng quan tâm định tính rộng, không chứa mâu thuẫn cơ học cụ thể, "
        "không có biến độc lập/phụ thuộc định lượng, và không thể bác bỏ (unfalsifiable). Do đó, nó chưa thể trở thành một đề tài nghiên cứu học thuật.",
        bold_prefix="Vấn đề tồn tại: "
    )
    add_p(
        "Sử dụng bằng chứng y văn đã được trích xuất và kiểm chứng trong repository để thu hẹp dần phạm vi nghiên cứu, "
        "loại bỏ các tuyên bố tính mới hình thức, chuyển dịch từ một chủ đề rộng sang một bài toán cơ học kết cấu cụ thể và một câu hỏi nghiên cứu có thể kiểm chứng và bác bỏ được.",
        bold_prefix="Mục tiêu thu hẹp: "
    )

    # -------------------------------------------------------------
    # SECTION 2 — TIMELINE THU HẸP PHẠM VI
    # -------------------------------------------------------------
    add_h1("SECTION 2 — TIMELINE THU HẸP PHẠM VI")
    add_p("Tiến trình thu hẹp tuân thủ chuỗi nhân quả: Tôi đã làm gì → Search keywords → Paper chính → Tôi rút ra gì → Vì sao/Evidence → Scope mới → Checkpoint.")

    tbl1 = doc.add_table(rows=8, cols=8)
    headers1 = ["Bước", "Tôi đã làm gì?", "Search keywords", "Paper chính", "Tôi rút ra gì?", "Vì sao / Evidence", "Scope mới", "Checkpoint"]
    for j, h in enumerate(headers1):
        tbl1.cell(0, j).paragraphs[0].text = h
    
    data1 = [
        ("1", "Khảo sát cơ chế biến đổi độ cứng cho robot mềm.",
         "Representative:\n'variable stiffness soft robotics'",
         "Narang et al. (2018)\nAdv. Funct. Mater.\n10.1002/adfm.201707136",
         "Layer jamming vượt trội về tỷ số thay đổi độ cứng uốn (tăng theo n²) và tốc độ đáp ứng chân không.",
         "Narang et al. (2018) chứng minh độ cứng uốn tăng gấp n² khi hút chân không.",
         "Khóa cơ chế kẹt lớp ma sát chân không (vacuum layer jamming). Loại bỏ cơ chế nhiệt, hạt kẹt.",
         "Cơ chế kẹt lớp được xác lập."),
        
        ("2", "Nhận diện hiện tượng vật lý chi phối sự suy giảm độ cứng khi dầm uốn.",
         "Representative:\n'layer jamming mechanics interlayer slip friction'",
         "Narang et al. (2018)\nCaruso et al. (2023)\n10.1016/j.ijmecsci.2023.108325",
         "Độ cứng giảm phi tuyến do hiện tượng trượt tương đối giữa các lớp (interlayer slip) bị khống chế bởi ma sát Coulomb τ ≤ μp.",
         "Narang (2018) và Caruso (2023) xác định 3 chế độ: tiền trượt, trượt một phần, trượt hoàn toàn.",
         "Tập trung vào cơ học tiếp xúc ma sát và trượt giao diện Coulomb. Loại bỏ giả định đàn hồi tuyến tính.",
         "Hiện tượng trượt ma sát giao diện được xác lập."),
        
        ("3", "Khảo sát mô hình giải tích dự đoán quan hệ lực - biến dạng.",
         "Historical:\nForward citations từ Narang et al. (2018)\nRepresentative:\n'layer jamming beam model analytical slip'",
         "Caruso et al. (2023)\nInt. J. Mech. Sci.\n10.1016/j.ijmecsci.2023.108325",
         "Mô hình giải tích rời rạc từng lớp (discrete layer-by-layer) cho n lớp chẵn đã tồn tại, nhưng số phương trình tăng theo n.",
         "Caruso et al. (2023) giải hệ phương trình giải tích rời rạc với thứ tự trượt từ tâm ra ngoài.",
         "Loại bỏ ý tưởng 'xây dựng mô hình giải tích đầu tiên'. Chuyển sang tìm kiếm mô hình rút gọn liên tục.",
         "Nhu cầu mô hình rút gọn continuum được xác lập."),
        
        ("4", "Tìm kiếm mô hình môi trường liên tục đồng hóa cho dầm layer jamming.",
         "Historical:\nForward citations từ Caruso et al. (2023)\nRepresentative:\n'continuum model layer jamming homogenized'",
         "Zhang et al. (2025)\nMech. Sci.\n10.5194/ms-16-821-2025",
         "Mô hình dầm liên tục (CLJM) đã được công bố hoàn chỉnh với biến trạng thái nửa chiều cao trượt ys(s).",
         "Zhang et al. (2025), Eqs. (3)–(9), pp. 823–825 công bố mô hình continuum đàn - dẻo lý tưởng.",
         "Loại bỏ hoàn toàn claim: 'xây dựng mô hình continuum đầu tiên'. Khóa mô hình Zhang 2025 làm mô hình cơ sở M1.",
         "Mô hình rút gọn M1 được lựa chọn."),
        
        ("5", "Phân định cấp độ mô hình: Dầm kết cấu vs. RVE vi mô.",
         "Historical:\nForward citations trên 10.5194/ms-16-821-2025\nRepresentative:\n'continuum modeling layer jamming RVE'",
         "Zhang et al. (2026)\nTheor. Appl. Mech. Lett.\n10.1016/j.taml.2025.100633",
         "Mô hình constitutive RVE vi mô dΣ = C^ep : dE chưa giải bài toán kết cấu dầm chịu tải.",
         "Zhang et al. (2026) chỉ mô hình hóa RVE 3D, không có nghiệm giải tích kết cấu và không có thực nghiệm.",
         "Loại bỏ hướng RVE vi mô. Khóa chọn mô hình dầm kết cấu M1 (Zhang 2025).",
         "Cấp độ mô hình kết cấu dầm được khóa."),
        
        ("6", "Deconstruct các giới hạn cơ học của mô hình Zhang 2025.",
         "Historical:\n'multi-leaf spring homogenization slip validity'\nRepresentative:\n'finite layer continuum comparison model error'",
         "Steif & Trojnacki (1993)\nMassabò & Campi (2014)\nWang et al. (2026, IJSS)\n10.1016/j.ijsolstr.2025.113689",
         "Mô hình continuum bỏ qua số lớp hữu hạn, áp suất pháp tự sinh; sai số lớn gần ngàm và khi uốn sâu.",
         "Massabò & Campi (2014) chỉ ra sai số mô hình đồng hóa 60–80% gần ngàm; Steif (1993) định nghĩa sai số tiệm cận.",
         "Loại bỏ ý tưởng 'kiểm chứng vài điểm thuận lợi'. Chuyển thành bài toán xác lập miền hiệu lực và biên gãy đổ.",
         "Bài toán ranh giới hiệu lực được định vị."),
        
        ("7", "Khóa hướng nghiên cứu D1 và đối soát công trình mới nhất 2026.",
         "Historical:\nD1-SEARCH-W10-001, 002\nRepresentative:\n'model validity domain layer jamming breakdown'",
         "Wang et al. (2026, M&D)\n10.1016/j.matdes.2026.116573\nYang et al. (2025, Sci. Rep.)",
         "Mô hình giải tích đơn giản bị sai lệch lớn trong thực nghiệm nhưng y văn chưa lập bản đồ biên hiệu lực định trước.",
         "Wang et al. (2026, M&D), tr. 15 xác nhận sai lệch lớn nhưng không có mô hình tham chiếu R và dung sai định trước.",
         "Loại bỏ mọi tuyên bố định tính; khóa bài toán lập bản đồ định lượng biên hiệu lực có kiểm chứng 2 phía biên.",
         "Hướng nghiên cứu D1 hoàn chỉnh được khóa.")
    ]
    for i, row in enumerate(data1):
        for j, val in enumerate(row):
            tbl1.cell(i+1, j).paragraphs[0].text = val
    col_w1 = [Cm(0.9), Cm(2.2), Cm(2.7), Cm(2.4), Cm(2.4), Cm(2.3), Cm(2.0), Cm(1.6)]
    format_table(tbl1, col_w1, font_size=12.0)

    # -------------------------------------------------------------
    # SECTION 3 — CÁC PAPER LÀM THAY ĐỔI SCOPE
    # -------------------------------------------------------------
    add_h1("SECTION 3 — CÁC PAPER LÀM THAY ĐỔI SCOPE")
    add_p("Bảng dưới đây tổng hợp các công trình then chốt trực tiếp làm thay đổi hoặc thu hẹp phạm vi nghiên cứu:")

    tbl3 = doc.add_table(rows=9, cols=5)
    headers3 = ["Paper (Tác giả, Năm, DOI)", "Paper cho thấy gì?", "Evidence cụ thể", "Tôi rút ra gì?", "D1 thu hẹp thế nào?"]
    for j, h in enumerate(headers3):
        tbl3.cell(0, j).paragraphs[0].text = h
    
    data3 = [
        ("Narang et al. (2018)\nAdv. Funct. Mater.\n10.1002/adfm.201707136",
         "Cơ chế kẹt lớp ma sát chân không; 3 chế độ trượt; mô hình 2 lớp giải tích và FEA tiếp xúc nhiều lớp.",
         "Công thức giải tích 2 lớp; độ cứng uốn tỷ lệ n²; FEA Abaqus tiếp xúc từng lớp.",
         "Layer jamming đã được chứng minh cơ học vững chắc; trượt ma sát Coulomb là cơ chế chi phối.",
         "Khóa cơ chế kẹt lớp; loại bỏ các cơ chế biến đổi độ cứng khác."),
        
        ("Caruso et al. (2023)\nInt. J. Mech. Sci.\n10.1016/j.ijmecsci.2023.108325",
         "Mô hình giải tích rời rạc từng lớp cho dầm n lớp chẵn; trượt tuần tự từ trục trung hòa ra ngoài.",
         "Hệ phương trình giải tích từng lớp; tính tổn hao năng lượng ma sát; thực nghiệm uốn 3 điểm.",
         "Mô hình giải tích rời rạc từng lớp đã tồn tại nhưng độ phức tạp giải tích tăng nhanh theo n.",
         "Loại bỏ ý tưởng 'xây dựng mô hình giải tích đầu tiên'; mở ra nhu cầu mô hình continuum."),
        
        ("Zhang, Yao, Zhao & Wei (2025)\nMech. Sci.\n10.5194/ms-16-821-2025",
         "Mô hình dầm liên tục (CLJM); biến trạng thái nửa chiều cao trượt ys(s); ngưỡng trượt Q_slip = (2/3)μpbh.",
         "Eqs. (3)–(9), pp. 823–825; giả thiết n → ∞; thuật toán gia số phi tuyến; sai lệch gần ngàm.",
         "Mô hình continuum dầm đã được công bố; dựa trên giả thiết n vô hạn và bỏ qua áp suất pháp tự sinh.",
         "Loại bỏ claim 'xây dựng mô hình continuum đầu tiên'; chọn Zhang 2025 làm M1; chuyển sang đánh giá hiệu lực."),
        
        ("Zhang, Yao, Li & Chen (2026)\nTheor. Appl. Mech. Lett.\n10.1016/j.taml.2025.100633",
         "Mô hình constitutive RVE vi mô dΣ = C^ep : dE từ 2 lớp tiếp xúc ma sát vi mô; kiểm chứng FEA 3D.",
         "Mô hình vật liệu vi mô 3D RVE; không có nghiệm giải tích dầm; tác giả nêu rõ thiếu thực nghiệm.",
         "Mô hình RVE vi mô là mô hình vật liệu, chưa giải bài toán kết cấu dầm thực tế.",
         "Loại bỏ hướng RVE vi mô; tập trung vào mô hình dầm kết cấu M1 (Zhang 2025)."),
        
        ("Steif & Trojnacki (1993a,b)\nJ. Appl. Mech.\n10.1115/1.2900989\n10.1115/1.2900990",
         "Sự hội tụ tiệm cận từ dầm nhiều lớp rời rạc trượt sang mô hình liên tục yếu trượt khi n → ∞.",
         "Part I (mô hình rời rạc) & Part II (mô hình liên tục); định nghĩa sai số định lượng e1, e2.",
         "So sánh discrete vs. continuum là bài toán kinh điển; mô hình liên tục luôn có sai số ở n hữu hạn.",
         "Cung cấp cơ sở lý thuyết đòi hỏi phải có chỉ số sai số định lượng tường minh cho D1."),
        
        ("Massabò & Campi (2014)\nMeccanica\n10.1007/s11012-014-9994-x",
         "Đánh giá sai số của các lý thuyết đồng hóa dầm/tấm nhiều lớp có trượt giao diện không hoàn hảo.",
         "Bỏ qua năng lượng biến dạng giao diện gây sai số mô hình 60–80%, tập trung nghiêm trọng gần biên ngàm.",
         "Giả thiết đồng hóa liên tục gãy đổ có hệ thống tại vùng biên ngàm và nơi biến thiên ứng suất lớn.",
         "Xác định nguyên nhân cơ học gây gãy đổ mô hình continuum; định hình bài toán biên hiệu lực."),
        
        ("Wang, Han, Yong & Zhou (2026)\nInt. J. Solids Struct.\n10.1016/j.ijsolstr.2025.113689",
         "Mô hình toàn cục - cục bộ cho kết cấu nhiều lớp; so sánh với FEA tiếp xúc trên dải n=10–150.",
         "Sử dụng ngưỡng sai số 5% để phân định ranh giới hiệu lực cho cuộn dây siêu dẫn; không có chân không.",
         "Việc dùng ngưỡng sai số phân định mô hình đã có tiền lệ, nhưng ở Wang 2026 là ngưỡng chọn hồi tố.",
         "Loại bỏ việc lấy ngưỡng 5% làm tính mới; buộc D1 phải xác lập dung sai dựa trên độ không đảm bảo đo."),
        
        ("Wang et al. (2026, M&D)\nMater. Des. 268, 116573\n10.1016/j.matdes.2026.116573",
         "Hệ dầm phao thổi khí gia cường multilayer jamming (MLJ); nhận định mô hình đơn giản sai lệch lớn.",
         "Section Discussion, p. 15: mô hình đơn giản dự đoán ~28 N, thấp hơn nhiều so với thực nghiệm.",
         "Sự sai lệch của mô hình đơn giản đã được quan sát, nhưng bài báo KHÔNG lập bản đồ biên hiệu lực.",
         "Audit disposition: PARTIAL_OVERLAP. Xác nhận khoảng trống D1: lập bản đồ biên định lượng có kiểm chứng 2 phía.")
    ]
    for i, row in enumerate(data3):
        for j, val in enumerate(row):
            tbl3.cell(i+1, j).paragraphs[0].text = val
    col_w3 = [Cm(3.3), Cm(3.4), Cm(3.3), Cm(3.2), Cm(3.3)]
    format_table(tbl3, col_w3, font_size=12.0)

    # -------------------------------------------------------------
    # SECTION 4 — NHỮNG CLAIM ĐÃ BỊ LOẠI
    # -------------------------------------------------------------
    add_h1("SECTION 4 — NHỮNG CLAIM ĐÃ BỊ LOẠI")
    add_p("Bảng dưới đây tổng hợp các tuyên bố ban đầu quá rộng đã bị loại bỏ dứt khoát sau khi đối chiếu với bằng chứng y văn:")

    tbl4 = doc.add_table(rows=7, cols=4)
    headers4 = ["Broad claim ban đầu", "Evidence phản bác", "Kết luận", "Scope thay thế"]
    for j, h in enumerate(headers4):
        tbl4.cell(0, j).paragraphs[0].text = h
    
    data4 = [
        ("Claim 1: 'Cơ chế kẹt lớp (layer jamming) chưa được mô hình hóa toán học đầy đủ.'",
         "Narang et al. (2018) đã giải 2 lớp; Caruso et al. (2023) đã thiết lập hệ giải tích n lớp chẵn.",
         "SAI. Mô hình hóa giải tích dầm layer jamming đã có tiền lệ vững chắc.",
         "Khảo sát tính hiệu lực và giới hạn sai số của một lớp mô hình cụ thể."),
        
        ("Claim 2: 'Hiện tượng trượt tương đối (interlayer slip) trong dầm kẹt lớp chưa được nghiên cứu.'",
         "Y văn dầm composite (partial interaction) và ma sát tiếp xúc đã nghiên cứu trượt giao diện hàng chục năm.",
         "SAI VÀ QUÁ RỘNG. Động học trượt ma sát Coulomb là cơ chế đã biết.",
         "Khảo sát trượt ma sát gắn liền với áp suất chân không và hiệu ứng chuyển thang liên tục."),
        
        ("Claim 3: 'Chưa có mô hình môi trường liên tục (continuum model) cho dầm layer jamming.'",
         "Zhang et al. (2025, Mech. Sci.) đã công bố mô hình continuum dầm CLJM.",
         "HOÀN TOÀN SAI. Mô hình dầm continuum đã tồn tại.",
         "Chọn mô hình Zhang 2025 làm mô hình cơ sở M1; chuyển sang đánh giá giới hạn hiệu lực."),
        
        ("Claim 4: 'Mô hình dầm layer jamming chưa từng được kiểm chứng thực nghiệm.'",
         "Narang (2018), Caruso (2023), Yang (2025), Wang (2026) đều đã làm thực nghiệm uốn 3 điểm / công-xôn.",
         "SAI. Thực nghiệm lực - độ võng thông thường là quy chuẩn cơ bản.",
         "Bố trí thực nghiệm có chủ đích kiểm chứng tại 2 phía của biên hiệu lực dự đoán."),
        
        ("Claim 5: 'Chưa ai so sánh mô hình rút gọn với mô hình số bậc cao (full FEA).'",
         "Narang (2018) so sánh analytical vs FEA; Wang (2026, IJSS) so sánh reduced vs contact FEA.",
         "QUÁ RỘNG. Việc so sánh hai cấp độ mô hình tự nó không đủ tính mới.",
         "So sánh M1 với mô hình tiếp xúc chi tiết R nhằm lượng hóa sai số dạng mô hình (model-form error)."),
        
        ("Claim 6: 'Chưa ai nghiên cứu sự sai lệch hoặc giới hạn của mô hình layer jamming.'",
         "Zhang (2025) nêu sai lệch gần ngàm; Wang (2026, M&D) chỉ ra sai lệch do hiệu ứng vỏ và gối tựa.",
         "QUÁ RỘNG. Sự tồn tại của sai lệch đã được ghi nhận định tính.",
         "Lập bản đồ ranh giới chấp nhận được định lượng dưới các tiêu chí dung sai định trước.")
    ]
    for i, row in enumerate(data4):
        for j, val in enumerate(row):
            tbl4.cell(i+1, j).paragraphs[0].text = val
    col_w4 = [Cm(3.9), Cm(4.5), Cm(3.6), Cm(4.5)]
    format_table(tbl4, col_w4, font_size=12.0)

    # -------------------------------------------------------------
    # SECTION 5 — 5 WHY
    # -------------------------------------------------------------
    add_h1("SECTION 5 — 5 WHY")
    add_p("Chuỗi 5-Why đào sâu bản chất cơ học, mỗi câu hỏi sản sinh truy vấn tìm kiếm và xác lập phạm vi nghiên cứu mới:")

    tbl5 = doc.add_table(rows=6, cols=5)
    headers5 = ["Why", "Câu hỏi", "Tôi rút ra gì?", "Evidence", "Search keywords tiếp theo"]
    for j, h in enumerate(headers5):
        tbl5.cell(0, j).paragraphs[0].text = h
    
    data5 = [
        ("Why 1", "Tại sao dầm layer jamming lại đáng nghiên cứu?",
         "Layer jamming cho dải biến đổi độ cứng uốn rất lớn (n²) và điều khiển nhanh bằng chân không.",
         "Narang et al. (2018), Adv. Funct. Mater.",
         "Representative:\n'layer jamming beam modeling analytical model'"),
        
        ("Why 2", "Tại sao không thiết kế thêm một cơ cấu robot mới?",
         "Y văn đã có hàng trăm cơ cấu; trở ngại cốt tử là thiếu công cụ mô hình hóa giải tích tin cậy.",
         "Caruso et al. (2023), Int. J. Mech. Sci.",
         "Historical:\nForward citations từ Caruso et al. (2023)\nRepresentative:\n'continuum model layer jamming beam'"),
        
        ("Why 3", "Tại sao chọn mô hình continuum thay vì mô hình rời rạc?",
         "Khi n lớn, mô hình rời rạc quá phức tạp, còn FEA tiếp xúc tốn kém tính toán. Mô hình continuum giải nhanh, phù hợp thiết kế.",
         "Zhang et al. (2025), Mech. Sci.",
         "Historical:\nForward citations trên 10.5194/ms-16-821-2025\nRepresentative:\n'continuum model layer jamming limitation'"),
        
        ("Why 4", "Tại sao tiếp tục nghiên cứu khi mô hình continuum đã có?",
         "Mô hình Zhang dựa trên giả thiết n→∞, bỏ qua áp suất pháp tự sinh; sai số lớn khi n nhỏ hoặc uốn sâu.",
         "Zhang (2025), Mech. Sci.; Wang (2026), Mater. Des.",
         "Historical:\n'multi-leaf spring homogenization slip validity'\nRepresentative:\n'model validity domain breakdown boundary'"),
        
        ("Why 5", "Tại sao xác định miền hiệu lực là bài toán đáng giá nhất?",
         "Kỹ sư không thể dùng mô hình mù quáng nếu không biết miền an toàn. Lập bản đồ định lượng VALID / INVALID tạo ra tri thức thiết kế nền tảng.",
         "Chưa có công trình nào lập bản đồ hiệu lực định trước cho mô hình continuum dầm kẹt lớp và kiểm chứng 2 phía biên.",
         "Khóa Hướng D1:\n'validity limits of homogenized slip model'")
    ]
    for i, row in enumerate(data5):
        for j, val in enumerate(row):
            tbl5.cell(i+1, j).paragraphs[0].text = val
    col_w5 = [Cm(1.2), Cm(3.8), Cm(4.4), Cm(3.4), Cm(3.7)]
    format_table(tbl5, col_w5, font_size=12.0)

    # -------------------------------------------------------------
    # SECTION 6 — TỪ TOPIC ĐẾN RESEARCH QUESTION
    # -------------------------------------------------------------
    add_h1("SECTION 6 — TỪ TOPIC ĐẾN RESEARCH QUESTION")
    add_p("Tiến trình chuyển dịch nhận thức từ một vùng quan tâm rộng đến câu hỏi nghiên cứu cụ thể:")

    p_flow = doc.add_paragraph()
    p_flow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_flow.paragraph_format.line_spacing = 1.15
    p_flow.paragraph_format.space_before = Pt(2)
    p_flow.paragraph_format.space_after = Pt(4)
    run_flow = p_flow.add_run(
        "Soft Robotics (Robot mềm)\n"
        "↓\n"
        "Variable Stiffness (Cơ chế biến đổi độ cứng thích ứng)\n"
        "↓\n"
        "Vacuum Layer Jamming (Kẹt lớp ma sát chân không)\n"
        "↓\n"
        "Layer-Jamming Beam Mechanics (Cơ học dầm kẹt lớp & trượt Coulomb)\n"
        "↓\n"
        "Continuum Modeling of Layer Jamming (Mô hình dầm liên tục CLJM của Zhang 2025)\n"
        "↓\n"
        "Limitations & Discrepancies of CLJM (Sai số do số lớp hữu hạn và uốn sâu)\n"
        "↓\n"
        "D1 Research Problem & Research Question"
    )
    run_flow.font.name = 'Times New Roman'
    run_flow.font.size = Pt(11)
    run_flow.font.bold = True
    run_flow.font.color.rgb = RGBColor(0, 32, 96)

    add_p(
        "Chưa có hiểu biết định lượng về miền hiệu lực (validity domain) và điều kiện gãy đổ cơ học (breakdown conditions) "
        "của mô hình continuum dầm Zhang et al. (2025) dưới tác động đồng thời của số lớp hữu hạn, áp suất chân không và mức độ uốn, "
        "khi đối sánh với mô hình tham chiếu tiếp xúc tường minh từng lớp và thực nghiệm vật lý.",
        bold_prefix="Current Research Problem: "
    )

    add_p(
        "\"Under what combinations of finite layer count, vacuum pressure, and bending severity does the Zhang et al. (2025) "
        "continuum layer-jamming beam model remain within predeclared output-specific model-form error tolerances relative to an explicit "
        "full-layer frictional-contact reference and physical experiments?\"",
        bold_prefix="Current Research Question — English: "
    )

    add_p(
        "\"Trong những tổ hợp nào của số lớp hữu hạn, áp suất chân không và mức độ uốn, mô hình continuum dầm layer-jamming của "
        "Zhang et al. (2025) duy trì được sai số dạng mô hình (model-form error) nằm trong các giới hạn dung sai chấp nhận định trước đối với "
        "từng đại lượng đầu ra cụ thể, khi so sánh với mô hình tham chiếu phần tử hữu hạn tiếp xúc ma sát tường minh từng lớp và thực nghiệm vật lý?\"",
        bold_prefix="Current Research Question — Vietnamese: "
    )

    # -------------------------------------------------------------
    # SECTION 7 — CHECKPOINT HIỆN TẠI
    # -------------------------------------------------------------
    add_h1("SECTION 7 — CHECKPOINT HIỆN TẠI")
    add_p("Trạng thái hiện tại của đề tài D1 sau các vòng rà soát y văn và kiểm chứng đối soát:")

    tbl7 = doc.add_table(rows=13, cols=2)
    headers7 = ["Checkpoint", "Status"]
    for j, h in enumerate(headers7):
        tbl7.cell(0, j).paragraphs[0].text = h
    
    data7 = [
        ("Broad field identified (Soft robotics / variable stiffness)", "DONE"),
        ("Variable-stiffness mechanism narrowed (Vacuum layer jamming)", "DONE"),
        ("Layer-jamming mechanics identified (Coulomb friction / interlayer slip)", "DONE"),
        ("Key existing models identified (Narang 2018, Caruso 2023, Zhang 2025)", "DONE"),
        ("Broad novelty claims removed (Model creation claims eliminated)", "DONE"),
        ("Specified baseline model selected (Zhang et al. 2025 CLJM as M1)", "DONE"),
        ("Major prior-art threats checked (Massabò 2014, Wang 2026 IJSS, Yang 2025)", "DONE"),
        ("Wang et al. 2026 checked (Full-text audit complete; partial overlap)", "DONE"),
        ("Candidate validity/breakdown problem established", "DONE"),
        ("Mentor-level research question formulated", "READY"),
        ("Exact numerical tolerances frozen (ε_w, ε_K, ε_Q)", "OPEN (Xin ý kiến Mentor)"),
        ("Final novelty adjudication", "NOT CLAIMED (Chờ thực thi M-R-E)")
    ]
    for i, row in enumerate(data7):
        for j, val in enumerate(row):
            tbl7.cell(i+1, j).paragraphs[0].text = val
    col_w7 = [Cm(12.2), Cm(4.3)]
    format_table(tbl7, col_w7, font_size=12.0)

    # -------------------------------------------------------------
    # SECTION 8 — REFERENCES
    # -------------------------------------------------------------
    add_h1("SECTION 8 — REFERENCES")
    add_p("Danh mục tài liệu tham khảo chuẩn mực APA 7th được kiểm chứng đối soát 100% với cơ sở dữ liệu bằng chứng repository:")

    refs = [
        "Caruso, M., Cosmi, F., Degrassi, A., Lanza, S., Totaro, M., & Beccai, L. (2023). Layer jamming: Modeling and experimental validation. International Journal of Mechanical Sciences, 252, 108325. https://doi.org/10.1016/j.ijmecsci.2023.108325",
        "Massabò, R., & Campi, F. (2014). Assessment and correction of theories for multilayered plates with imperfect interfaces. Meccanica, 49(9), 2151–2169. https://doi.org/10.1007/s11012-014-9994-x",
        "Narang, Y. S., Vlassak, J. J., & Howe, R. D. (2018). Mechanically versatile soft machines through laminar jamming. Advanced Functional Materials, 28(17), 1707136. https://doi.org/10.1002/adfm.201707136",
        "Steif, P. S., & Trojnacki, A. (1993a). Bending stress enhancement in materials with limited shear resistance—Part I. Slipping-layers model. Journal of Applied Mechanics, 60(4), 834–839. https://doi.org/10.1115/1.2900989",
        "Steif, P. S., & Trojnacki, A. (1993b). Bending stress enhancement in materials with limited shear resistance—Part II. Unlayered shear-weak model. Journal of Applied Mechanics, 60(4), 840–846. https://doi.org/10.1115/1.2900990",
        "Wang, Q., Feng, P., Wu, B., Teng, M., Jansen, K., & Bao, C. (2026). Multilayer jamming-reinforced inflatable systems for rapidly deployable lightweight construction. Materials & Design, 268, 116573. https://doi.org/10.1016/j.matdes.2026.116573",
        "Wang, S., Han, Y., Yong, H., & Zhou, Y. (2026). The global-local mechanical behaviors of multilayered structure and applications to superconducting coils. International Journal of Solids and Structures, 311, 113689. https://doi.org/10.1016/j.ijsolstr.2025.113689",
        "Yang, X., Guo, S., & Wang, H. (2025). Behavior of layer jamming plate with tunable stiffness and its near wake structure in cross flow. Scientific Reports, 15, 22364. https://doi.org/10.1038/s41598-025-22364-w",
        "Zhang, S., Yao, J., Zhao, W., & Wei, C. (2025). A continuum-based model for a layer jamming beam. Mechanical Sciences, 16(2), 821–830. https://doi.org/10.5194/ms-16-821-2025",
        "Zhang, S., Yao, J., Li, S., & Chen, X. (2026). Continuum modeling for layer jamming structures. Theoretical and Applied Mechanics Letters, 16(1), 100633. https://doi.org/10.1016/j.taml.2025.100633"
    ]
    
    for ref in refs:
        p_ref = doc.add_paragraph()
        p_ref.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_ref.paragraph_format.line_spacing = 1.15
        p_ref.paragraph_format.space_after = Pt(3)
        p_ref.paragraph_format.left_indent = Inches(0.4)
        p_ref.paragraph_format.first_line_indent = Inches(-0.4)
        r = p_ref.add_run(ref)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11.0)

    out_path = "docs/reports/D1_MENTOR_RESEARCH_NARROWING_REPORT.docx"
    doc.save(out_path)
    print(f"SUCCESS: Generated {out_path}")

def generate_markdown_report():
    md_content = """# BÁO CÁO KHOA HỌC: TIẾN TRÌNH THU HẸP PHẠM VI NGHIÊN CỨU
## TỪ CHỦ ĐỀ TỔNG QUAN ĐẾN CÂU HỎI NGHIÊN CỨU D1
### (GIỚI HẠN HIỆU LỰC MÔ HÌNH CONTINUUM DẦM LAYER JAMMING)

**Báo cáo tiến trình rà soát y văn trình Cán bộ hướng dẫn (Mentor Progress Report)**  
**Chuyên ngành:** Cơ kỹ thuật / Robot mềm | **Thời điểm:** 26/09/2026  
**Phương pháp luận:** Falsification-First & Evidence-Based Narrowing  

---

# SECTION 1 — MỤC TIÊU BAN ĐẦU

* **Chủ đề ban đầu (Initial Topic):** Robot mềm (*soft robotics*) và bài toán điều biến độ cứng thích ứng (*variable stiffness*).
* **Vấn đề tồn tại:** Khái niệm "nghiên cứu dầm layer jamming" là một vùng quan tâm định tính rộng, không chứa mâu thuẫn cơ học cụ thể, không có biến độc lập/phụ thuộc định lượng, và không thể bác bỏ (*unfalsifiable*). Do đó, nó chưa thể trở thành một đề tài nghiên cứu học thuật.
* **Mục tiêu thu hẹp:** Sử dụng bằng chứng y văn đã được trích xuất và kiểm chứng trong repository để thu hẹp dần phạm vi nghiên cứu, loại bỏ các tuyên bố tính mới hình thức, chuyển dịch từ một chủ đề rộng sang một bài toán cơ học kết cấu cụ thể và một câu hỏi nghiên cứu có thể kiểm chứng và bác bỏ được.

---

# SECTION 2 — TIMELINE THU HẸP PHẠM VI

Chuỗi nhân quả tối thiểu: Hành động ($X$) $\\rightarrow$ Truy vấn ($Y$) $\\rightarrow$ Bài báo ($Z$) $\\rightarrow$ Bằng chứng ($A$) $\\rightarrow$ Kết luận & Loại bỏ ($B, C$) $\\rightarrow$ Scope mới ($D$) $\\rightarrow$ Checkpoint.

| Bước | Tôi đã làm gì? | Search keywords | Paper chính | Tôi rút ra gì? | Vì sao / Evidence | Scope mới | Checkpoint |
| :---: | :--- | :--- | :--- | :--- | :--- | :--- | :--- |
| **1** | Khảo sát cơ chế biến đổi độ cứng cho robot mềm. | **Representative:**<br>`variable stiffness soft robotics` | Narang et al. (2018)<br>*Adv. Funct. Mater.*<br>[10.1002/adfm.201707136](https://doi.org/10.1002/adfm.201707136) | Layer jamming vượt trội về tỷ số thay đổi độ cứng uốn ($n^2$) và tốc độ đáp ứng chân không. | Narang et al. (2018) chứng minh độ cứng uốn tăng gấp $n^2$ khi hút chân không. | Khóa cơ chế kẹt lớp ma sát chân không (*vacuum layer jamming*). Loại bỏ cơ chế nhiệt, hạt kẹt. | Cơ chế kẹt lớp được xác lập. |
| **2** | Nhận diện hiện tượng vật lý chi phối sự suy giảm độ cứng khi dầm uốn. | **Representative:**<br>`layer jamming mechanics interlayer slip friction` | Narang et al. (2018)<br>Caruso et al. (2023)<br>[10.1016/j.ijmecsci.2023.108325](https://doi.org/10.1016/j.ijmecsci.2023.108325) | Độ cứng giảm phi tuyến do hiện tượng trượt tương đối giữa các lớp (*interlayer slip*) bị khống chế bởi ma sát Coulomb $\\tau \\le \\mu p$. | Narang (2018) và Caruso (2023) xác định 3 chế độ: tiền trượt, trượt một phần, trượt hoàn toàn. | Tập trung vào cơ học tiếp xúc ma sát và trượt giao diện Coulomb. Loại bỏ giả định đàn hồi tuyến tính. | Hiện tượng trượt ma sát giao diện được xác lập. |
| **3** | Khảo sát mô hình giải tích dự đoán quan hệ lực - biến dạng. | **Historical:**<br>Forward citations từ Narang et al. (2018)<br>**Representative:**<br>`layer jamming beam model analytical slip` | Caruso et al. (2023)<br>*Int. J. Mech. Sci.*<br>[10.1016/j.ijmecsci.2023.108325](https://doi.org/10.1016/j.ijmecsci.2023.108325) | Mô hình giải tích rời rạc từng lớp (*discrete layer-by-layer*) cho $n$ lớp chẵn đã tồn tại, nhưng số phương trình tăng theo $n$. | Caruso et al. (2023) giải hệ phương trình giải tích rời rạc với thứ tự trượt từ tâm ra ngoài. | Loại bỏ ý tưởng "xây dựng mô hình giải tích đầu tiên". Chuyển sang tìm kiếm mô hình rút gọn liên tục. | Nhu cầu mô hình rút gọn continuum được xác lập. |
| **4** | Tìm kiếm mô hình môi trường liên tục đồng hóa cho dầm layer jamming. | **Historical:**<br>Forward citations từ Caruso et al. (2023)<br>**Representative:**<br>`continuum model layer jamming homogenized` | Zhang et al. (2025)<br>*Mech. Sci.*<br>[10.5194/ms-16-821-2025](https://doi.org/10.5194/ms-16-821-2025) | Mô hình dầm liên tục (CLJM) đã được công bố hoàn chỉnh với biến trạng thái nửa chiều cao trượt $y_s(s)$. | Zhang et al. (2025), Eqs. (3)–(9), pp. 823–825 công bố mô hình continuum đàn - dẻo lý tưởng. | Loại bỏ hoàn toàn claim: "xây dựng mô hình continuum đầu tiên". Khóa mô hình Zhang 2025 làm mô hình cơ sở $M_1$. | Mô hình rút gọn $M_1$ được lựa chọn. |
| **5** | Phân định cấp độ mô hình: Dầm kết cấu vs. RVE vi mô. | **Historical:**<br>Forward citations trên `10.5194/ms-16-821-2025`<br>**Representative:**<br>`continuum modeling layer jamming RVE` | Zhang et al. (2026)<br>*Theor. Appl. Mech. Lett.*<br>[10.1016/j.taml.2025.100633](https://doi.org/10.1016/j.taml.2025.100633) | Mô hình constitutive RVE vi mô $d\\Sigma = C^{ep} : dE$ chưa giải bài toán kết cấu dầm chịu tải. | Zhang et al. (2026) chỉ mô hình hóa RVE 3D, không có nghiệm giải tích kết cấu và không có thực nghiệm. | Loại bỏ hướng RVE vi mô. Khóa chọn mô hình dầm kết cấu $M_1$ (Zhang 2025). | Cấp độ mô hình kết cấu dầm được khóa. |
| **6** | Deconstruct các giới hạn cơ học của mô hình Zhang 2025. | **Historical:**<br>`multi-leaf spring homogenization slip validity`<br>**Representative:**<br>`finite layer continuum comparison model error` | Steif & Trojnacki (1993)<br>Massabò & Campi (2014)<br>Wang et al. (2026, IJSS)<br>[10.1016/j.ijsolstr.2025.113689](https://doi.org/10.1016/j.ijsolstr.2025.113689) | Mô hình continuum bỏ qua số lớp hữu hạn, áp suất pháp tự sinh; sai số lớn gần ngàm và khi uốn sâu. | Massabò & Campi (2014) chỉ ra sai số mô hình đồng hóa 60–80% gần ngàm; Steif (1993) định nghĩa sai số tiệm cận. | Loại bỏ ý tưởng "kiểm chứng vài điểm thuận lợi". Chuyển thành bài toán xác lập miền hiệu lực và biên gãy đổ. | Bài toán ranh giới hiệu lực được định vị. |
| **7** | Khóa hướng nghiên cứu D1 và đối soát công trình mới nhất 2026. | **Historical:**<br>D1-SEARCH-W10-001, 002<br>**Representative:**<br>`model validity domain layer jamming breakdown` | Wang et al. (2026, M&D)<br>[10.1016/j.matdes.2026.116573](https://doi.org/10.1016/j.matdes.2026.116573)<br>Yang et al. (2025, Sci. Rep.) | Mô hình giải tích đơn giản bị sai lệch lớn trong thực nghiệm nhưng y văn chưa lập bản đồ biên hiệu lực định trước. | Wang et al. (2026, M&D), tr. 15 xác nhận sai lệch lớn nhưng không có mô hình tham chiếu $R$ và dung sai định trước. | Loại bỏ mọi tuyên bố định tính; khóa bài toán lập bản đồ định lượng biên hiệu lực có kiểm chứng 2 phía biên. | Hướng nghiên cứu D1 hoàn chỉnh được khóa. |

---

# SECTION 3 — CÁC PAPER LÀM THAY ĐỔI SCOPE

| Paper (Tác giả, Năm, DOI) | Paper cho thấy gì? | Evidence cụ thể | Tôi rút ra gì? | D1 thu hẹp thế nào? |
| :--- | :--- | :--- | :--- | :--- |
| **Narang et al. (2018)**<br>*Adv. Funct. Mater.*<br>[10.1002/adfm.201707136](https://doi.org/10.1002/adfm.201707136) | Cơ chế kẹt lớp ma sát chân không; 3 chế độ trượt; mô hình 2 lớp giải tích và FEA tiếp xúc nhiều lớp. | Công thức giải tích 2 lớp; độ cứng uốn tỷ lệ $n^2$; FEA Abaqus tiếp xúc từng lớp. | Layer jamming đã được chứng minh cơ học vững chắc; trượt ma sát Coulomb là cơ chế chi phối. | Khóa cơ chế kẹt lớp; loại bỏ các cơ chế biến đổi độ cứng khác. |
| **Caruso et al. (2023)**<br>*Int. J. Mech. Sci.*<br>[10.1016/j.ijmecsci.2023.108325](https://doi.org/10.1016/j.ijmecsci.2023.108325) | Mô hình giải tích rời rạc từng lớp cho dầm $n$ lớp chẵn; trượt tuần tự từ trục trung hòa ra ngoài. | Hệ phương trình giải tích từng lớp; tính tổn hao năng lượng ma sát; thực nghiệm uốn 3 điểm. | Mô hình giải tích rời rạc từng lớp đã tồn tại nhưng độ phức tạp giải tích tăng nhanh theo $n$. | Loại bỏ ý tưởng "xây dựng mô hình giải tích đầu tiên"; mở ra nhu cầu mô hình continuum. |
| **Zhang, Yao, Zhao & Wei (2025)**<br>*Mech. Sci.*<br>[10.5194/ms-16-821-2025](https://doi.org/10.5194/ms-16-821-2025) | Mô hình dầm liên tục (CLJM); biến trạng thái nửa chiều cao trượt $y_s(s)$; ngưỡng trượt $Q_{slip} = \\frac{2}{3}\\mu p b h$. | Eqs. (3)–(9), pp. 823–825; giả thiết $n \\to \\infty$; thuật toán gia số phi tuyến; sai lệch gần ngàm. | Mô hình continuum dầm đã được công bố; dựa trên giả thiết $n$ vô hạn và bỏ qua áp suất pháp tự sinh. | Loại bỏ claim "xây dựng mô hình continuum đầu tiên"; chọn Zhang 2025 làm $M_1$; chuyển sang đánh giá hiệu lực. |
| **Zhang, Yao, Li & Chen (2026)**<br>*Theor. Appl. Mech. Lett.*<br>[10.1016/j.taml.2025.100633](https://doi.org/10.1016/j.taml.2025.100633) | Mô hình constitutive RVE vi mô $d\\Sigma = C^{ep} : dE$ từ 2 lớp tiếp xúc ma sát vi mô; kiểm chứng FEA 3D. | Mô hình vật liệu vi mô 3D RVE; không có nghiệm giải tích dầm; tác giả nêu rõ thiếu thực nghiệm. | Mô hình RVE vi mô là mô hình vật liệu, chưa giải bài toán kết cấu dầm thực tế. | Loại bỏ hướng RVE vi mô; tập trung vào mô hình dầm kết cấu $M_1$ (Zhang 2025). |
| **Steif & Trojnacki (1993a,b)**<br>*J. Appl. Mech.*<br>[10.1115/1.2900989](https://doi.org/10.1115/1.2900989)<br>[10.1115/1.2900990](https://doi.org/10.1115/1.2900990) | Sự hội tụ tiệm cận từ dầm nhiều lớp rời rạc trượt sang mô hình liên tục yếu trượt khi $n \\to \\infty$. | Part I (mô hình rời rạc) & Part II (mô hình liên tục); định nghĩa sai số định lượng $e_1, e_2$. | So sánh discrete vs. continuum là bài toán kinh điển; mô hình liên tục luôn có sai số ở $n$ hữu hạn. | Cung cấp cơ sở lý thuyết đòi hỏi phải có chỉ số sai số định lượng tường minh cho D1. |
| **Massabò & Campi (2014)**<br>*Meccanica*<br>[10.1007/s11012-014-9994-x](https://doi.org/10.1007/s11012-014-9994-x) | Đánh giá sai số của các lý thuyết đồng hóa dầm/tấm nhiều lớp có trượt giao diện không hoàn hảo. | Bỏ qua năng lượng biến dạng giao diện gây sai số mô hình 60–80%, tập trung nghiêm trọng gần biên ngàm. | Giả thiết đồng hóa liên tục gãy đổ có hệ thống tại vùng biên ngàm và nơi biến thiên ứng suất lớn. | Xác định nguyên nhân cơ học gây gãy đổ mô hình continuum; định hình bài toán biên hiệu lực. |
| **Wang, Han, Yong & Zhou (2026)**<br>*Int. J. Solids Struct.*<br>[10.1016/j.ijsolstr.2025.113689](https://doi.org/10.1016/j.ijsolstr.2025.113689) | Mô hình toàn cục - cục bộ cho kết cấu nhiều lớp; so sánh với FEA tiếp xúc trên dải $n=10–150$. | Sử dụng ngưỡng sai số 5% để phân định ranh giới hiệu lực cho cuộn dây siêu dẫn; không có chân không. | Việc dùng ngưỡng sai số phân định mô hình đã có tiền lệ, nhưng ở Wang 2026 là ngưỡng chọn hồi tố. | Loại bỏ việc lấy ngưỡng 5% làm tính mới; buộc D1 phải xác lập dung sai dựa trên độ không đảm bảo đo. |
| **Wang et al. (2026, M&D)**<br>*Mater. Des.* 268, 116573<br>[10.1016/j.matdes.2026.116573](https://doi.org/10.1016/j.matdes.2026.116573) | Hệ dầm phao thổi khí gia cường multilayer jamming (MLJ); nhận định mô hình đơn giản sai lệch lớn. | Section Discussion, p. 15: mô hình đơn giản dự đoán ~28 N, thấp hơn nhiều so với thực nghiệm. | Sự sai lệch của mô hình đơn giản đã được quan sát, nhưng bài báo KHÔNG lập bản đồ biên hiệu lực. | Audit disposition: PARTIAL_OVERLAP. Xác nhận khoảng trống D1: lập bản đồ biên định lượng có kiểm chứng 2 phía. |

---

# SECTION 4 — NHỮNG CLAIM ĐÃ BỊ LOẠI

| Broad claim ban đầu | Evidence phản bác | Kết luận | Scope thay thế |
| :--- | :--- | :--- | :--- |
| **Claim 1:** "Cơ chế kẹt lớp (layer jamming) chưa được mô hình hóa toán học đầy đủ." | Narang et al. (2018) đã giải 2 lớp; Caruso et al. (2023) đã thiết lập hệ giải tích $n$ lớp chẵn. | **SAI.** Mô hình hóa giải tích dầm layer jamming đã có tiền lệ vững chắc. | Khảo sát tính hiệu lực và giới hạn sai số của một lớp mô hình cụ thể. |
| **Claim 2:** "Hiện tượng trượt tương đối (interlayer slip) trong dầm kẹt lớp chưa được nghiên cứu." | Y văn dầm composite (*partial interaction*) và ma sát tiếp xúc đã nghiên cứu trượt giao diện hàng chục năm. | **SAI VÀ QUÁ RỘNG.** Động học trượt ma sát Coulomb là cơ chế đã biết. | Khảo sát trượt ma sát gắn liền với áp suất chân không và hiệu ứng chuyển thang liên tục. |
| **Claim 3:** "Chưa có mô hình môi trường liên tục (continuum model) cho dầm layer jamming." | Zhang et al. (2025, *Mech. Sci.*) đã công bố mô hình continuum dầm CLJM. | **HOÀN TOÀN SAI.** Mô hình dầm continuum đã tồn tại. | Chọn mô hình Zhang 2025 làm mô hình cơ sở $M_1$; chuyển sang đánh giá giới hạn hiệu lực. |
| **Claim 4:** "Mô hình dầm layer jamming chưa từng được kiểm chứng thực nghiệm." | Narang (2018), Caruso (2023), Yang (2025), Wang (2026) đều đã làm thực nghiệm uốn 3 điểm / công-xôn. | **SAI.** Thực nghiệm lực - độ võng thông thường là quy chuẩn cơ bản. | Bố trí thực nghiệm có chủ đích kiểm chứng tại 2 phía của biên hiệu lực dự đoán. |
| **Claim 5:** "Chưa ai so sánh mô hình rút gọn với mô hình số bậc cao (full FEA)." | Narang (2018) so sánh analytical vs FEA; Wang (2026, IJSS) so sánh reduced vs contact FEA. | **QUÁ RỘNG.** Việc so sánh hai cấp độ mô hình tự nó không đủ tính mới. | So sánh $M_1$ với mô hình tiếp xúc chi tiết $R$ nhằm lượng hóa sai số dạng mô hình (*model-form error*). |
| **Claim 6:** "Chưa ai nghiên cứu sự sai lệch hoặc giới hạn của mô hình layer jamming." | Zhang (2025) nêu sai lệch gần ngàm; Wang (2026, M&D) chỉ ra sai lệch do hiệu ứng vỏ và gối tựa. | **QUÁ RỘNG.** Sự tồn tại của sai lệch đã được ghi nhận định tính. | Lập bản đồ ranh giới chấp nhận được định lượng dưới các tiêu chí dung sai định trước. |

---

# SECTION 5 — 5 WHY

| Why | Câu hỏi | Tôi rút ra gì? | Evidence | Search keywords tiếp theo |
| :---: | :--- | :--- | :--- | :--- |
| **Why 1** | Tại sao dầm layer jamming lại đáng nghiên cứu? | Layer jamming cho dải biến đổi độ cứng uốn rất lớn ($n^2$) và điều khiển nhanh bằng chân không. | Narang et al. (2018), *Adv. Funct. Mater.* | **Representative:**<br>`layer jamming beam modeling analytical model` |
| **Why 2** | Tại sao không thiết kế thêm một cơ cấu robot mới? | Y văn đã có hàng trăm cơ cấu; trở ngại cốt tử là thiếu công cụ mô hình hóa giải tích tin cậy. | Caruso et al. (2023), *Int. J. Mech. Sci.* | **Historical:**<br>Forward citations từ Caruso et al. (2023)<br>**Representative:**<br>`continuum model layer jamming beam` |
| **Why 3** | Tại sao chọn mô hình continuum thay vì mô hình rời rạc? | Khi $n$ lớn, mô hình rời rạc quá phức tạp, còn FEA tiếp xúc tốn kém tính toán. Mô hình continuum giải nhanh, phù hợp thiết kế. | Zhang et al. (2025), *Mech. Sci.* | **Historical:**<br>Forward citations trên `10.5194/ms-16-821-2025`<br>**Representative:**<br>`continuum model layer jamming limitation` |
| **Why 4** | Tại sao tiếp tục nghiên cứu khi mô hình continuum đã có? | Mô hình Zhang dựa trên giả thiết $n \\to \\infty$, bỏ qua áp suất pháp tự sinh; sai số lớn khi $n$ nhỏ hoặc uốn sâu. | Zhang (2025), *Mech. Sci.*; Wang (2026), *Mater. Des.* | **Historical:**<br>`multi-leaf spring homogenization slip validity`<br>**Representative:**<br>`model validity domain breakdown boundary` |
| **Why 5** | Tại sao xác định miền hiệu lực là bài toán đáng giá nhất? | Kỹ sư không thể dùng mô hình mù quáng nếu không biết miền an toàn. Lập bản đồ định lượng VALID / INVALID tạo ra tri thức thiết kế nền tảng. | Chưa có công trình nào lập bản đồ hiệu lực định trước cho mô hình continuum dầm kẹt lớp và kiểm chứng 2 phía biên. | **Khóa Hướng D1:**<br>`validity limits of homogenized slip model` |

---

# SECTION 6 — TỪ TOPIC ĐẾN RESEARCH QUESTION

```text
Soft Robotics (Robot mềm)
↓
Variable Stiffness (Cơ chế biến đổi độ cứng thích ứng)
↓
Vacuum Layer Jamming (Kẹt lớp ma sát chân không)
↓
Layer-Jamming Beam Mechanics (Cơ học dầm kẹt lớp & trượt Coulomb)
↓
Continuum Modeling of Layer Jamming (Mô hình dầm liên tục CLJM của Zhang 2025)
↓
Limitations & Discrepancies of CLJM (Sai số do số lớp hữu hạn và uốn sâu)
↓
D1 Research Problem & Research Question
```

* **Current Research Problem:**  
  Chưa có hiểu biết định lượng về miền hiệu lực (*validity domain*) và điều kiện gãy đổ cơ học (*breakdown conditions*) của mô hình continuum dầm Zhang et al. (2025) dưới tác động đồng thời của số lớp hữu hạn, áp suất chân không và mức độ uốn, khi đối sánh với mô hình tham chiếu tiếp xúc tường minh từng lớp và thực nghiệm vật lý.

* **Current Research Question — English:**  
  *"Under what combinations of finite layer count, vacuum pressure, and bending severity does the Zhang et al. (2025) continuum layer-jamming beam model remain within predeclared output-specific model-form error tolerances relative to an explicit full-layer frictional-contact reference and physical experiments?"*

* **Current Research Question — Vietnamese:**  
  *"Trong những tổ hợp nào của số lớp hữu hạn, áp suất chân không và mức độ uốn, mô hình continuum dầm layer-jamming của Zhang et al. (2025) duy trì được sai số dạng mô hình (model-form error) nằm trong các giới hạn dung sai chấp nhận định trước đối với từng đại lượng đầu ra cụ thể, khi so sánh với mô hình tham chiếu phần tử hữu hạn tiếp xúc ma sát tường minh từng lớp và thực nghiệm vật lý?"*

---

# SECTION 7 — CHECKPOINT HIỆN TẠI

| Checkpoint | Status |
| :--- | :---: |
| Broad field identified (Soft robotics / variable stiffness) | **DONE** |
| Variable-stiffness mechanism narrowed (Vacuum layer jamming) | **DONE** |
| Layer-jamming mechanics identified (Coulomb friction / interlayer slip) | **DONE** |
| Key existing models identified (Narang 2018, Caruso 2023, Zhang 2025) | **DONE** |
| Broad novelty claims removed (Model creation claims eliminated) | **DONE** |
| Specified baseline model selected (Zhang et al. 2025 CLJM as $M_1$) | **DONE** |
| Major prior-art threats checked (Massabò 2014, Wang 2026 IJSS, Yang 2025) | **DONE** |
| Wang et al. 2026 checked (Full-text audit complete; partial overlap) | **DONE** |
| Candidate validity/breakdown problem established | **DONE** |
| Mentor-level research question formulated | **READY** |
| Exact numerical tolerances frozen ($\\varepsilon_w, \\varepsilon_K, \\varepsilon_Q$) | **OPEN (Xin ý kiến Mentor)** |
| Final novelty adjudication | **NOT CLAIMED (Chờ thực thi M-R-E)** |

---

# SECTION 8 — REFERENCES

1. Caruso, M., Cosmi, F., Degrassi, A., Lanza, S., Totaro, M., & Beccai, L. (2023). Layer jamming: Modeling and experimental validation. *International Journal of Mechanical Sciences*, 252, 108325. https://doi.org/10.1016/j.ijmecsci.2023.108325
2. Massabò, R., & Campi, F. (2014). Assessment and correction of theories for multilayered plates with imperfect interfaces. *Meccanica*, 49(9), 2151–2169. https://doi.org/10.1007/s11012-014-9994-x
3. Narang, Y. S., Vlassak, J. J., & Howe, R. D. (2018). Mechanically versatile soft machines through laminar jamming. *Advanced Functional Materials*, 28(17), 1707136. https://doi.org/10.1002/adfm.201707136
4. Steif, P. S., & Trojnacki, A. (1993a). Bending stress enhancement in materials with limited shear resistance—Part I. Slipping-layers model. *Journal of Applied Mechanics*, 60(4), 834–839. https://doi.org/10.1115/1.2900989
5. Steif, P. S., & Trojnacki, A. (1993b). Bending stress enhancement in materials with limited shear resistance—Part II. Unlayered shear-weak model. *Journal of Applied Mechanics*, 60(4), 840–846. https://doi.org/10.1115/1.2900990
6. Wang, Q., Feng, P., Wu, B., Teng, M., Jansen, K., & Bao, C. (2026). Multilayer jamming-reinforced inflatable systems for rapidly deployable lightweight construction. *Materials & Design*, 268, 116573. https://doi.org/10.1016/j.matdes.2026.116573
7. Wang, S., Han, Y., Yong, H., & Zhou, Y. (2026). The global-local mechanical behaviors of multilayered structure and applications to superconducting coils. *International Journal of Solids and Structures*, 311, 113689. https://doi.org/10.1016/j.ijsolstr.2025.113689
8. Yang, X., Guo, S., & Wang, H. (2025). Behavior of layer jamming plate with tunable stiffness and its near wake structure in cross flow. *Scientific Reports*, 15, 22364. https://doi.org/10.1038/s41598-025-22364-w
9. Zhang, S., Yao, J., Zhao, W., & Wei, C. (2025). A continuum-based model for a layer jamming beam. *Mechanical Sciences*, 16(2), 821–830. https://doi.org/10.5194/ms-16-821-2025
10. Zhang, S., Yao, J., Li, S., & Chen, X. (2026). Continuum modeling for layer jamming structures. *Theoretical and Applied Mechanics Letters*, 16(1), 100633. https://doi.org/10.1016/j.taml.2025.100633
"""
    with open("docs/reports/D1_MENTOR_RESEARCH_NARROWING_REPORT.md", "w", encoding="utf-8") as f:
        f.write(md_content)
    print("SUCCESS: Generated docs/reports/D1_MENTOR_RESEARCH_NARROWING_REPORT.md")

if __name__ == '__main__':
    generate_markdown_report()
    build_report()
