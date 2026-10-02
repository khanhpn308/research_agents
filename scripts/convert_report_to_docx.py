import re
import os
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_hyperlink(paragraph, url, text):
    clean_url = url.strip('<> ')
    try:
        part = paragraph.part
        r_id = part.relate_to(clean_url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        hyperlink = OxmlElement('w:hyperlink')
        hyperlink.set(qn('r:id'), r_id)
        new_run = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), 'Times New Roman')
        rFonts.set(qn('w:hAnsi'), 'Times New Roman')
        rPr.append(rFonts)

        # Gạch chân, không màu
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        rPr.append(u)

        new_run.append(rPr)
        text_node = OxmlElement('w:t')
        text_node.text = text
        new_run.append(text_node)
        hyperlink.append(new_run)
        paragraph._p.append(hyperlink)
    except Exception:
        run = paragraph.add_run(text)
        run.font.name = 'Times New Roman'
        run.underline = True

def parse_inline_into_paragraph(paragraph, text, base_font_size=11, default_italic=False):
    token_pattern = re.compile(
        r'(\*{0,2}\[(?:HỒ SƠ QUYẾT ĐỊNH|VERIFIED FULL TEXT|METADATA ONLY|INFERENCE[^\]]*|HYPOTHESIS[^\]]*)\]\*{0,2})'
        r'|(\[[^\]]+\]\((?:<[^>]+>|[^\)]+)\))'
        r'|(`[^`]+`)'
        r'|(\*\*\*[^*]+\*\*\*)'
        r'|(\*\*[^*]+\*\*)'
        r'|(\*[^*]+\*)'
    )

    pos = 0
    for match in token_pattern.finditer(text):
        start, end = match.span()
        if start > pos:
            plain_text = text[pos:start]
            run = paragraph.add_run(plain_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(base_font_size)
            if default_italic:
                run.italic = True

        badge_match = match.group(1)
        link_match = match.group(2)
        code_match = match.group(3)
        bold_italic_match = match.group(4)
        bold_match = match.group(5)
        italic_match = match.group(6)

        if badge_match:
            clean_badge = badge_match.strip('*')
            run = paragraph.add_run(clean_badge)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(base_font_size)
            run.bold = True
            # BỎ MÀU: Không gán color
        elif link_match:
            m = re.match(r'\[([^\]]+)\]\((<[^>]+>|[^\)]+)\)', link_match)
            if m:
                link_text, link_url = m.groups()
                link_url = link_url.strip('<> ')
                add_hyperlink(paragraph, link_url, link_text)
            else:
                run = paragraph.add_run(link_match)
                run.font.name = 'Times New Roman'
                run.font.size = Pt(base_font_size)
        elif code_match:
            c_text = code_match[1:-1]
            run = paragraph.add_run(c_text)
            run.font.name = 'Consolas'
            run.font.size = Pt(base_font_size - 0.5)
            # BỎ MÀU
        elif bold_italic_match:
            bi_text = bold_italic_match[3:-3]
            run = paragraph.add_run(bi_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(base_font_size)
            run.bold = True
            run.italic = True
        elif bold_match:
            b_text = bold_match[2:-2]
            run = paragraph.add_run(b_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(base_font_size)
            run.bold = True
            if default_italic:
                run.italic = True
        elif italic_match:
            i_text = italic_match[1:-1]
            run = paragraph.add_run(i_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(base_font_size)
            run.italic = True

        pos = end

    if pos < len(text):
        plain_text = text[pos:]
        run = paragraph.add_run(plain_text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(base_font_size)
        if default_italic:
            run.italic = True

def get_table_col_widths(headers, num_cols):
    h_str = " ".join(headers).lower()
    if 'hướng' in h_str and 'candidate' in h_str:
        return [0.65, 2.15, 2.05, 1.82]
    elif 'bài' in h_str and 'suy yếu' in h_str:
        return [1.80, 1.90, 1.45, 1.52]
    elif 'nguồn' in h_str and 'doi' in h_str:
        return [2.10, 1.50, 3.07]
    elif 'vòng' in h_str and 'verdict' in h_str:
        return [0.95, 2.05, 1.50, 2.17]
    elif 'thành phần' in h_str and 'vai trò' in h_str:
        return [0.85, 2.85, 2.97]
    elif 'id' in h_str and 'claim/câu hỏi trước' in h_str:
        return [0.55, 1.50, 1.60, 1.45, 1.57]
    elif 'cụm trong tên' in h_str:
        return [1.70, 2.50, 2.47]
    elif 'sai khác' in h_str and 'xử lý' in h_str:
        return [2.90, 3.77]
    
    total_width = 6.67
    if num_cols == 2:
        return [2.5, 4.17]
    elif num_cols == 3:
        return [1.8, 2.3, 2.57]
    elif num_cols == 4:
        return [1.3, 1.8, 1.8, 1.77]
    elif num_cols == 5:
        return [0.65, 1.45, 1.55, 1.45, 1.57]
    else:
        w = total_width / num_cols
        return [w] * num_cols

def convert_markdown_strict_plain(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')
    doc = docx.Document()

    # Layout A4, margins 1 inch
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        # KHÔNG HEADER:
        header = section.header
        header.is_linked_to_previous = False
        for p in header.paragraphs:
            p.text = ""

        # KHÔNG FOOTER:
        footer = section.footer
        footer.is_linked_to_previous = False
        for p in footer.paragraphs:
            p.text = ""

    # Normal style: Times New Roman 12pt, không gán color (để Word tự động)
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)

    i = 0
    n = len(lines)

    while i < n:
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue

        # Title (# ...) -> IN ĐẬM, KHÔNG MÀU, KHÔNG KẺ DÒNG
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            # Tuyệt đối không thêm pBdr (không kẻ dòng phân cách)

            title_text = line[2:].strip()
            run = p.add_run(title_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(16)
            run.bold = True
            # Không gán color
            i += 1
            continue

        # Heading 2 (## ...) -> IN ĐẬM, KHÔNG MÀU, KHÔNG KẺ DÒNG
        if line.startswith('## '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True

            h2_text = line[3:].strip()
            run = p.add_run(h2_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(13)
            run.bold = True
            # Không gán color
            i += 1
            continue

        # Heading 3 (### ...) -> IN ĐẬM, KHÔNG MÀU, KHÔNG KẺ DÒNG
        if line.startswith('### '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True

            h3_text = line[4:].strip()
            run = p.add_run(h3_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.bold = True
            # Không gán color
            i += 1
            continue

        # Blockquote (> ...) -> CHỈ THỤT LỀ, KHÔNG KHUNG, KHÔNG KẺ DÒNG, KHÔNG MÀU
        if line.startswith('>'):
            quote_lines = []
            while i < n and (lines[i].startswith('>') or not lines[i].strip()):
                if lines[i].startswith('>'):
                    q_line = lines[i][1:].strip()
                    quote_lines.append(q_line)
                elif not lines[i].strip():
                    if i + 1 < n and lines[i+1].startswith('>'):
                        quote_lines.append('')
                    else:
                        break
                i += 1

            for ql in quote_lines:
                if not ql:
                    continue
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.4)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                parse_inline_into_paragraph(p, ql, base_font_size=11, default_italic=True)
            continue

        # Table (| ... |) -> BẢNG LƯỚI ĐƠN GIẢN, KHÔNG ĐÁNH MÀU
        if line.startswith('|') and '|' in line[1:]:
            table_lines = []
            while i < n and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            rows_data = []
            for tline in table_lines:
                content_inner = tline.strip()
                if content_inner.startswith('|'):
                    content_inner = content_inner[1:]
                if content_inner.endswith('|'):
                    content_inner = content_inner[:-1]
                cols = [c.strip() for c in content_inner.split('|')]
                if all(re.match(r'^:?-+:?$', c) for c in cols):
                    continue
                rows_data.append(cols)

            if rows_data:
                num_rows = len(rows_data)
                num_cols = max(len(r) for r in rows_data)
                
                for r in rows_data:
                    while len(r) < num_cols:
                        r.append('')

                table = doc.add_table(rows=num_rows, cols=num_cols)
                table.style = 'Table Grid'
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = False

                header_tr = table.rows[0]._tr.get_or_add_trPr()
                header_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
                for row in table.rows:
                    trPr = row._tr.get_or_add_trPr()
                    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

                col_widths = get_table_col_widths(rows_data[0], num_cols)
                scale = 6.27 / 6.67
                col_widths = [w * scale for w in col_widths]
                col_widths_dxa = [int(w * 1440) for w in col_widths]

                tblGrid = table._tbl.tblGrid
                if tblGrid is not None:
                    tblGrid.clear()
                    for w_dxa in col_widths_dxa:
                        gridCol = parse_xml(f'<w:gridCol {nsdecls("w")} w:w="{w_dxa}"/>')
                        tblGrid.append(gridCol)

                for r_idx, row_content in enumerate(rows_data):
                    is_header = (r_idx == 0)
                    for c_idx, cell_value in enumerate(row_content):
                        cell = table.cell(r_idx, c_idx)
                        cell.width = Inches(col_widths[c_idx])
                        
                        tcPr = cell._tc.get_or_add_tcPr()
                        tcW = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{col_widths_dxa[c_idx]}" w:type="dxa"/>')
                        tcPr.append(tcW)

                        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)

                        cp = cell.paragraphs[0]
                        cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        cp.paragraph_format.space_before = Pt(1)
                        cp.paragraph_format.space_after = Pt(1)
                        cp.paragraph_format.line_spacing = 1.05

                        if is_header:
                            # In đậm, không màu
                            run = cp.add_run(cell_value)
                            run.font.name = 'Times New Roman'
                            run.font.size = Pt(10)
                            run.bold = True
                        else:
                            parse_inline_into_paragraph(cp, cell_value, base_font_size=10)

                sp = doc.add_paragraph()
                sp.paragraph_format.space_before = Pt(0)
                sp.paragraph_format.space_after = Pt(4)
            continue

        # Bullet list item (- ...)
        if line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            list_text = line[2:].strip()
            parse_inline_into_paragraph(p, list_text, base_font_size=11)
            i += 1
            continue

        # Numbered list item (1. ..., 2. ...)
        num_match = re.match(r'^(\d+)\.\s+(.*)$', line)
        if num_match:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            list_text = num_match.group(2)
            parse_inline_into_paragraph(p, list_text, base_font_size=11)
            i += 1
            continue

        # Subtitle / metadata line at the start
        if line.startswith('Ngày dựng lại:'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(8)
            parse_inline_into_paragraph(p, line, base_font_size=10, default_italic=True)
            i += 1
            continue

        # Standard Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        parse_inline_into_paragraph(p, line, base_font_size=11)
        i += 1

    doc.save(docx_path)
    print(f"Successfully generated strict plain docx: {docx_path}")

if __name__ == '__main__':
    md_file = "docs/reports/D1_DECISION_HISTORY_AND_EVIDENCE_FOR_SUPERVISOR_VI.md"
    docx_file = "docs/reports/D1_DECISION_HISTORY_AND_EVIDENCE_FOR_SUPERVISOR_VI.docx"
    convert_markdown_strict_plain(md_file, docx_file)
