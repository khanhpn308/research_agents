import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="B0B0B0", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def add_section_header(doc, text):
    p = doc.add_paragraph(style='List Paragraph')
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(13)
    run.bold = True
    return p

def add_subsection_header(doc, text):
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    run.bold = True
    return p

def add_body_paragraph(doc, text, italic=False):
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(11)
    if italic:
        run.italic = True
    return p

def add_styled_table(doc, headers, data_rows, col_widths=None):
    num_rows = len(data_rows) + 1
    num_cols = len(headers)
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Set table width 100%
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    tblW.set(qn('w:w'), '5000')
    tblW.set(qn('w:type'), 'pct')

    set_table_borders(table, color="C0C0C0", sz="4", val="single")

    # Header row
    hdr_row = table.rows[0]
    for c_idx, h_text in enumerate(headers):
        cell = hdr_row.cells[c_idx]
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        set_cell_shading(cell, "F2F4F7")
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(1)
        cp.paragraph_format.space_after = Pt(1)
        run = cp.add_run(h_text)
        run.font.name = 'Arial'
        run.font.size = Pt(10)
        run.bold = True

    # Data rows
    for r_idx, row_data in enumerate(data_rows):
        row = table.rows[r_idx + 1]
        for c_idx, cell_value in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            if r_idx % 2 == 1:
                set_cell_shading(cell, "FAFAFB")
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(1)
            cp.paragraph_format.space_after = Pt(1)
            run = cp.add_run(cell_value)
            run.font.name = 'Arial'
            run.font.size = Pt(9.5)

    if col_widths and len(col_widths) == num_cols:
        for row in table.rows:
            for c_idx, w in enumerate(col_widths):
                row.cells[c_idx].width = Inches(w)

    sp = doc.add_paragraph(style='Normal')
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)
    return table

def append_framework_to_doc(docx_path):
    doc = docx.Document(docx_path)

    # Add divider / main section header
    add_section_header(doc, "  Framework trích xuất nghiên cứu từ Evidence JSON (Phần 4 — Phần 11)")
    add_body_paragraph(doc, "Dưới đây là đặc tả chi tiết khung trích xuất nghiên cứu (Research Extraction Framework) từ tài liệu hướng dẫn evidence JSON (docs/learning/EVIDENCE_JSON_FRAMEWORK_GUIDE_VI.md), bao gồm các nhóm trường từ Phần 4 đến Phần 11 phục vụ trích xuất có cấu trúc, kiểm tra đối chứng và xây dựng ma trận bằng chứng khoa học cho dự án:")

    # 4. Thông tin trích xuất — model và usage
    add_subsection_header(doc, "4. Thông tin trích xuất — model và usage")
    sec4_headers = ["Trường", "Ý nghĩa"]
    sec4_rows = [
        ["model", "Model AI được ghi nhận đã thực hiện trích xuất"],
        ["usage.prompt_tokens", "Số token đầu vào được báo cáo"],
        ["usage.completion_tokens", "Số token đầu ra được provider báo cáo"],
        ["usage.total_tokens", "Tổng token được báo cáo"],
        ["usage.thinking_tokens", "Token reasoning nếu provider có cung cấp"],
        ["usage.cache_read_tokens", "Token đầu vào được đọc từ cache"]
    ]
    add_styled_table(doc, sec4_headers, sec4_rows, col_widths=[2.0, 4.5])
    add_body_paragraph(doc, "Lưu ý phương pháp: Đây là thông tin vận hành, không phải bằng chứng bài đúng hoặc đề tài mới. Các bộ đếm có thể phụ thuộc provider; không tự cộng lại các trường như thể tất cả đều là các phần độc lập.", italic=True)

    # 5. Thông tin thư mục trong paper
    add_subsection_header(doc, "5. Thông tin thư mục trong paper")
    sec5_headers = ["Trường", "Ý nghĩa", "Kiểu dữ liệu quan sát được"]
    sec5_rows = [
        ["title", "Tên bài", "Chuỗi"],
        ["authors", "Danh sách tác giả", "Danh sách chuỗi"],
        ["year", "Năm công bố được trích xuất", "Chuỗi, ví dụ \"2026\""],
        ["doi", "Định danh DOI", "Chuỗi"]
    ]
    add_styled_table(doc, sec5_headers, sec5_rows, col_widths=[1.5, 3.5, 1.5])
    add_body_paragraph(doc, "Lưu ý phương pháp: Cần phân biệt online-first và năm issue nếu khác nhau. DOI để trống không đồng nghĩa chắc chắn bài không có DOI. Không dùng năm trong filename thay metadata đã xác minh.", italic=True)

    # 6. Bài nghiên cứu vấn đề gì?
    add_subsection_header(doc, "6. Bài nghiên cứu vấn đề gì?")
    sec6_headers = ["Trường", "Ý nghĩa", "Câu hỏi giúp trả lời"]
    sec6_rows = [
        ["research_problem", "Vấn đề hoặc hạn chế bài muốn xử lý", "Vì sao cần nghiên cứu?"],
        ["research_objective", "Mục tiêu cụ thể của bài", "Tác giả định làm gì?"],
        ["robot_type", "Loại robot hoặc cấu trúc ứng dụng", "Hệ được nghiên cứu hoặc ứng dụng vào đâu?"],
        ["stiffness_mechanism", "Cơ chế tạo hoặc thay đổi độ cứng", "Vì sao độ cứng thay đổi?"],
        ["actuation", "Cách tác động/điều khiển hệ", "Hệ được kích hoạt bằng gì?"]
    ]
    add_styled_table(doc, sec6_headers, sec6_rows, col_widths=[1.8, 3.0, 1.7])
    add_body_paragraph(doc, "Lưu ý phương pháp: Ba trường đầu là chuỗi; stiffness_mechanism và actuation là danh sách trong dữ liệu đã kiểm tra. Cơ chế độ cứng khác cơ cấu tác động. Ví dụ minh họa: khí nén có thể là cách tác động, còn contact–friction là cơ chế khiến độ cứng thay đổi. Không coi hai trường này là từ đồng nghĩa.", italic=True)

    # 7. Mô hình và giả thiết
    add_subsection_header(doc, "7. Mô hình và giả thiết")
    sec7_headers = ["Trường", "Ý nghĩa"]
    sec7_rows = [
        ["modeling_methods", "Danh sách phương pháp mô hình hóa: analytical beam model, FEM, mô hình hiện tượng học…"],
        ["constitutive_assumptions", "Danh sách giả thiết về ứng xử vật liệu/cơ học được ghi nhận: đàn hồi tuyến tính, luật NiTi, Coulomb friction…"]
    ]
    add_styled_table(doc, sec7_headers, sec7_rows, col_widths=[2.2, 4.3])
    add_body_paragraph(doc, "Lưu ý phương pháp: Dữ liệu hiện có đôi khi đưa cả giả thiết hình học hoặc beam theory vào constitutive_assumptions. Tên trường không bảo đảm mọi mục đều là giả thiết cấu thành vật liệu theo nghĩa hẹp. Khi phân tích chuyên sâu, phải tách material law, kinematics, contact law và boundary conditions theo nội dung nguồn thực tế.", italic=True)

    # 8. Ba nhóm biến cần phân biệt
    add_subsection_header(doc, "8. Ba nhóm biến cần phân biệt")
    sec8_headers = ["Trường", "Ý nghĩa", "Ví dụ minh họa cho phép thử bó dây"]
    sec8_rows = [
        ["independent_variables", "Các biến chủ động thay đổi", "Áp suất, mức uốn"],
        ["dependent_variables", "Các đại lượng đáp ứng được đo/tính", "Lực, độ cứng, trượt tương đối"],
        ["control_variables", "Các yếu tố giữ cố định hoặc kiểm soát để so sánh công bằng", "Vật liệu dây, chiều dài mẫu, nhiệt độ, tốc độ tải"]
    ]
    add_styled_table(doc, sec8_headers, sec8_rows, col_widths=[2.0, 2.7, 1.8])
    add_body_paragraph(doc, "Lưu ý phương pháp: Cả ba trường là danh sách. Ví dụ trên chỉ minh họa phân loại, không phải protocol đã được thực hiện. Cùng một đại lượng có thể là independent variable trong nghiên cứu này nhưng là control variable trong nghiên cứu khác.", italic=True)

    # 9. Thí nghiệm và kết quả
    add_subsection_header(doc, "9. Thí nghiệm và kết quả")
    sec9_headers = ["Trường", "Ý nghĩa"]
    sec9_rows = [
        ["experimental_setup", "Mẫu, apparatus, cảm biến, cách gá và cách tạo tải"],
        ["performance_metrics", "Tiêu chí dùng đánh giá: độ cứng, sai số dự đoán, năng lượng tiêu tán…"],
        ["main_results", "Những kết quả chính được trích xuất từ bài"]
    ]
    add_styled_table(doc, sec9_headers, sec9_rows, col_widths=[2.2, 4.3])
    add_body_paragraph(doc, "Lưu ý phương pháp: Cả ba trường là danh sách. Phân biệt đo đại lượng gì (dependent_variables) với dùng đại lượng nào để đánh giá (performance_metrics). Không suy từ một kết quả nằm trong main_results rằng nó đã được independent experimental validation: cần kiểm tra kết quả là analytical, numerical, fitted hay measured, và dữ liệu nào đã được dùng để calibration.", italic=True)

    # 10. Bằng chứng liên quan — evidence_relevant_to_topic
    add_subsection_header(doc, "10. Bằng chứng liên quan — evidence_relevant_to_topic")
    add_body_paragraph(doc, "Đây là danh sách các phát biểu được chọn vì liên quan đến câu hỏi nghiên cứu của dự án. Mỗi mục có ba trường con:")
    sec10_headers = ["Trường con", "Ý nghĩa"]
    sec10_rows = [
        ["claim", "Phát biểu cụ thể được rút từ nguồn"],
        ["evidence_type", "Loại hỗ trợ cho phát biểu: analytical, numerical, experimental, background hoặc mô tả kết hợp"],
        ["page_numbers", "Danh sách trang được ghi để truy ngược PDF"]
    ]
    add_styled_table(doc, sec10_headers, sec10_rows, col_widths=[2.0, 4.5])
    add_body_paragraph(doc, "Lưu ý phương pháp: Một mục nằm trong trường này chưa chắc là kết quả thực nghiệm trực tiếp. Nó có thể là background mà bài đang dẫn từ nguồn khác. Cần kiểm tra loại bằng chứng và trang nguồn trước khi coi đó là primary evidence. page_numbers không mặc nhiên phân biệt printed page và số trang PDF. Nếu chuẩn bị citation chính thức, phải đối chiếu hệ đánh số. Không tự tạo equation/figure/section number khi JSON không lưu.", italic=True)

    # 11. Giới hạn, suy luận và hướng tiếp theo
    add_subsection_header(doc, "11. Giới hạn, suy luận và hướng tiếp theo")
    sec11_headers = ["Trường", "Ý nghĩa", "Cảnh báo"]
    sec11_rows = [
        ["limitations_stated_by_authors", "Giới hạn được ghi là do tác giả trực tiếp nêu", "Cần kiểm tra đúng wording và scope trong bài"],
        ["limitations_inferred", "Giới hạn do người/model phân tích suy ra", "Không gán thành kết luận của tác giả"],
        ["future_work", "Hướng nghiên cứu tiếp theo được trích xuất", "Future work năm trước không chứng minh vấn đề vẫn mở hiện nay"],
        ["possible_gap_implications", "Suy luận về ảnh hưởng của bài đối với gap dự án", "Không phải bằng chứng xác nhận novelty"],
        ["confidence", "Mức tự đánh giá độ tin cậy của bản trích xuất", "Không phải confidence đề tài mới; không bảo đảm trích xuất không sai"]
    ]
    add_styled_table(doc, sec11_headers, sec11_rows, col_widths=[2.2, 2.5, 1.8])
    add_body_paragraph(doc, "Lưu ý phương pháp: Bốn trường đầu là danh sách, confidence là chuỗi trong dữ liệu đã kiểm tra. Không phải mọi mục trong chúng đều có page locator riêng. Không biến limitations_inferred hoặc possible_gap_implications thành 'khoảng trống đã được chứng minh' khi chưa qua đối chiếu nguồn và adversarial audit.", italic=True)

    doc.save(docx_path)
    print(f"Successfully appended framework (Sections 4-11) to: {docx_path}")

if __name__ == '__main__':
    import sys
    target_path = sys.argv[1] if len(sys.argv) > 1 else 'docs/reports/MP1_DECISION_HISTORY_AND_EVIDENCE_FOR_SUPERVISOR_VI.docx'
    append_framework_to_doc(target_path)
